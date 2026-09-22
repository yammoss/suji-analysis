# -*- coding: utf-8 -*-
# 웹(Pyodide) 전용 연결부. 원본 도구 코드는 손대지 않고 여기서만 감싼다.
#   APP   = 도구 폴더 (/work/app) - 공유폴더에서 읽은 엑셀도 여기에 같은 이름으로 둔다
#   CACHE = 브라우저 저장소(IndexedDB)에 붙인 폴더 - 비용파일로 만든 기준데이터를 보관
import io
import json
import os
import pickle
import runpy
import sys
from pathlib import Path

APP = Path("/work/app")
CACHE = Path("/cache")
sys.path.insert(0, str(APP))

_ACT = {}


def _sig(p: Path) -> str:
    st = p.stat()
    return f"{st.st_size}_{int(st.st_mtime)}"


def _patch_actuals():
    """과거실적은 페이지를 연 동안 한 번만 읽는다 (파일이 바뀌면 다시)."""
    import profit_tool.actuals as A
    orig = getattr(A, "_web_orig", A.Actuals)
    A._web_orig = orig

    class Cached(orig):
        def __new__(cls, path=A.DEFAULT_ACTUALS):
            p = Path(path)
            key = (str(p), _sig(p)) if p.exists() else (str(p), "")
            if key not in _ACT:
                _ACT.clear()
                _ACT[key] = orig(path)
            return _ACT[key]

    A.Actuals = Cached


def status() -> str:
    """폴더에 무엇이 있는지 + 선택지 (JSON)."""
    import build_dataset as B
    cost = B.find_cost_file(APP)
    out = dict(cost=cost.name if cost else None,
               actuals=(APP / "과거실적 DATA.xlsx").exists(),
               calendar=(APP / "공휴일 DATA.xlsx").exists(),
               plan_tabs=[], daily=False)
    if out["actuals"]:
        import openpyxl
        from profit_tool.actuals import HEADER_ALIASES, is_plan_sheet
        wb = openpyxl.load_workbook(APP / "과거실적 DATA.xlsx", read_only=True)
        out["plan_tabs"] = [n for n in wb.sheetnames[1:] if is_plan_sheet(n)]
        head = next(wb.worksheets[0].iter_rows(max_row=1, values_only=True), ())
        wb.close()
        names = {n.lower() for n in HEADER_ALIASES["day"]}
        out["daily"] = out["calendar"] and any(str(h or "").strip().lower() in names for h in head)
    return json.dumps(out, ensure_ascii=False)


def ensure_dataset() -> str:
    """비용파일 → data/w26_dataset.json. 같은 파일로 만든 적 있으면 브라우저 보관본을 쓴다."""
    import build_dataset as B
    cost = B.find_cost_file(APP)
    if cost is None:
        raise RuntimeError("비용 추정용 파일(비용파일)을 폴더에서 찾지 못했습니다.")
    key = f"dataset_{_sig(cost)}.json"
    target = APP / "data" / "w26_dataset.json"
    target.parent.mkdir(parents=True, exist_ok=True)
    cached = CACHE / key
    if cached.exists():
        target.write_bytes(cached.read_bytes())
        return f"기준데이터 : 보관본 사용 ({cost.name})"
    _run_main(B.main, ["build_dataset.py"])
    for old in CACHE.glob("dataset_*.json"):
        old.unlink()
    cached.write_bytes(target.read_bytes())
    return f"기준데이터 : {cost.name} 에서 새로 만듦"


def _run_main(fn, argv):
    old = sys.argv
    sys.argv = list(argv)
    try:
        fn()
    except SystemExit as e:
        if e.code not in (None, 0):
            raise RuntimeError(str(e.code))
    finally:
        sys.argv = old


def run_script(script: str, argv: list) -> str:
    """도구 스크립트 실행. 새로 생긴 output 엑셀 경로(없으면 '')."""
    out_dir = APP / "output"
    out_dir.mkdir(exist_ok=True)
    before = {p: p.stat().st_mtime for p in out_dir.glob("*.xlsx")}
    _patch_actuals()
    os.chdir(APP)
    old = sys.argv
    sys.argv = [script] + list(argv)
    try:
        runpy.run_path(str(APP / script), run_name="__main__")
    except SystemExit as e:
        if e.code not in (None, 0):
            print(f"\n[중단] {e.code}")
    finally:
        sys.argv = old
    new = [p for p in out_dir.glob("*.xlsx") if before.get(p) != p.stat().st_mtime]
    return str(max(new, key=lambda p: p.stat().st_mtime)) if new else ""


def detect_period(text: str) -> str:
    from profit_tool.parser import pop_period
    try:
        found, _ = pop_period(text)
    except Exception:
        found = ""
    return found or ""
