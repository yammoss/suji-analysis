# -*- coding: utf-8 -*-
from __future__ import annotations
import json
import re
import sys
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as ci
BASE = Path(__file__).resolve().parent
OUT = BASE / 'data' / 'w26_dataset.json'
REQUIRED_SHEETS = ['기준데이터', '기종별 노선별 CASK', '사업량', '수입', '부대수입']
ALLOC_SUFFIX, DIRECT_SUFFIX = ('배부비', '노선별 직접비')
ALLOC_SHEET = DIRECT_SHEET = ''
COST_SHEET_MARKERS = {'기준데이터', '기종별 노선별 CASK', '사업량'}

def list_cost_files(folder: Path) -> list:
    out = []
    for f in folder.glob('*.xlsx'):
        if f.name.startswith('~$'):
            continue
        try:
            wb = openpyxl.load_workbook(f, read_only=True)
            names = set(wb.sheetnames)
            wb.close()
        except Exception:
            continue
        if COST_SHEET_MARKERS <= names:
            out.append(f)
    return sorted(out, key=lambda x: -x.stat().st_mtime)

def find_cost_file(folder: Path) -> Path | None:
    cands = list_cost_files(folder)
    return cands[0] if cands else None

def detect_months(wb) -> list[str]:
    plan = wb['사업량']
    months = []
    for c in range(ci('G'), ci('G') + 12):
        v = str(plan.cell(4, c).value or '')
        m = re.match('^\\s*(\\d{2,4})\\.(\\d{1,2})\\s*월\\s*$', v)
        if not m:
            break
        y = int(m.group(1))
        months.append(f'{(y if y < 100 else y % 100):02d}.{int(m.group(2)):02d}')
    if not months:
        raise ValueError("'사업량' 시트 G4 부터의 월 헤더('26.10월' 형식)를 찾지 못했습니다.")
    anc = wb['부대수입']
    n_anc = 0
    for c in range(ci('D'), ci('D') + 12):
        if re.match('^\\s*\\d{1,2}\\s*월\\s*$', str(anc.cell(27, c).value or '')):
            n_anc += 1
        else:
            break
    if n_anc:
        months = months[:n_anc]
    alloc = wb[next((s for s in wb.sheetnames if s.endswith(ALLOC_SUFFIX)))]
    ins = find_label_row(alloc, '항공기보험료') or 8
    last = 0
    for i in range(len(months)):
        v = alloc.cell(ins, ci('C') + i).value
        if isinstance(v, (int, float)) and v:
            last = i + 1
    return months[:last] if last else months
COL = {'사업량_편수': 'G', '사업량_공급석': 'AJ', '사업량_수송석': 'AX', '사업량_ASK': 'BM', '사업량_BT': 'CB', '수입_운송수입': 'AI', '배부비_pool': 'C', '부대수입_rask': 'D'}
POOL_LABELS = {'승객보상비': ['승객보상비'], '항공기보험료': ['항공기보험료'], '항공기임차료': ['항공기임차료'], '판매수수료': ['판매수수료'], '가맹점수수료': ['가맹점수수료'], '전산비': ['전산비']}
DEP_SECTION = '기종 배부용'
INDIRECT_SECTION = '간접비 배부용'
INDIRECT_BTFC_HINT = '승무원'
POOL_ROW: dict = {}
DEP_ROW: dict = {}
ESCALATION_FUEL_ROWCOL = (3, 'H')
ESCALATION_OTHER_ROWCOL = (3, 'C')
LEASE_EXCLUDED_TYPES = {'A332', 'B77W'}
CASK_FUEL_COL = 'I'
CASK_FX_COLS = ['J', 'L', 'N', 'P', 'S', 'V', 'X', 'Y', 'Z', 'AA']
CASK_KRW_COLS = ['K', 'M', 'O', 'Q', 'R', 'T', 'U', 'W']
CASK_ITEM_NAMES = {'I': '항공유', 'J': '영공통과료', 'K': '조업비', 'L': '조업비_해외', 'M': '공항시설사용료', 'N': '공항시설사용료_해외', 'O': '항공기착륙료', 'P': '항공기착륙료_해외', 'Q': '기내서비스비', 'R': '기내식비', 'S': '기내식비_해외', 'T': '비행수당', 'U': '착륙수당', 'V': '승무원퍼디엄', 'W': '레이오버비용', 'X': '레이오버비용_해외', 'Y': '항공기임차료_추가(MR)', 'Z': '정비비', 'AA': '정비비(Pooling)'}

def num(v, default=0.0):
    return float(v) if isinstance(v, (int, float)) else default

def _norm(v) -> str:
    return re.sub('[\\s,()·]', '', str(v or '')).upper()

def find_label_row(ws, label, col='B', end=200, exact_first=True):
    want = _norm(label)
    partial = None
    for r in range(1, end + 1):
        got = _norm(ws.cell(r, ci(col)).value)
        if not got:
            continue
        if got == want:
            return r
        if partial is None and (want in got or got in want):
            partial = r
    return None if exact_first else partial or None

def find_section_row(ws, text, col='A', end=200):
    want = _norm(text)
    for r in range(1, end + 1):
        if want in _norm(ws.cell(r, ci(col)).value):
            return r
    return None

def resolve_alloc_rows(alloc, aircraft_codes) -> list:
    warn = []
    POOL_ROW.clear()
    DEP_ROW.clear()
    for key, labels in POOL_LABELS.items():
        row = None
        for lab in labels:
            row = find_label_row(alloc, lab)
            if row:
                break
        if row is None:
            raise ValueError(f"배부비 시트에서 '{key}' 행을 찾지 못했습니다.\n   B열 라벨이 바뀌었다면 build_dataset.py 의 POOL_LABELS 를 맞춰주세요.")
        POOL_ROW[key] = row
    head = find_section_row(alloc, DEP_SECTION)
    if head is None:
        raise ValueError(f"배부비 시트에서 '{DEP_SECTION}' 섹션을 찾지 못했습니다.")
    want = {'ALL'} | set(aircraft_codes)
    for r in range(head + 1, head + 40):
        lab = str(alloc.cell(r, ci('B')).value or '').strip().upper()
        if lab in want:
            DEP_ROW[lab] = r
        elif lab.startswith('○') or _norm(alloc.cell(r, ci('A')).value).startswith('○'):
            break
    missing = want - set(DEP_ROW)
    if missing:
        warn.append(f"기종 감가상각 행을 못 찾은 기종: {', '.join(sorted(missing))} (해당 기종은 기종별 감가 배부 없이 계산됩니다)")
    head = find_section_row(alloc, INDIRECT_SECTION)
    if head is None:
        raise ValueError(f"배부비 시트에서 '{INDIRECT_SECTION}' 섹션을 찾지 못했습니다.")
    found = []
    for r in range(head + 1, head + 12):
        lab = str(alloc.cell(r, ci('B')).value or '').strip()
        if lab:
            found.append((r, lab))
    if len(found) < 2:
        raise ValueError(f"'{INDIRECT_SECTION}' 아래에서 간접고정비 2개 행을 찾지 못했습니다.")
    btfc = next((r for r, lab in found if INDIRECT_BTFC_HINT in lab), None)
    if btfc is None:
        btfc = found[1][0]
        warn.append(f"간접고정비 중 BT·편수 배부 행을 라벨('{INDIRECT_BTFC_HINT}')로 구분하지 못해 두 번째 행({btfc})을 썼습니다. 확인 필요.")
    rev = next((r for r, _ in found if r != btfc), None)
    POOL_ROW['간접고정_BTFC'] = btfc
    POOL_ROW['간접고정_운송수입'] = rev
    return warn
SUBTOTAL_LABELS = ('직접 변동비', '직접(배부비)', '간접(배부비)')

def audit_alloc_sheet(alloc, months_n) -> list:
    c0 = ci(COL['배부비_pool'])

    def tot(r):
        return sum((num(alloc.cell(r, c0 + i).value) for i in range(months_n)))
    rows = []
    for r in range(1, 120):
        lab = str(alloc.cell(r, ci('B')).value or '').strip()
        head = str(alloc.cell(r, ci('A')).value or '').strip()
        rows.append((r, lab, head, tot(r)))
    out, i = ([], 0)
    while i < len(rows):
        r, lab, head, v = rows[i]
        if not any((k.replace(' ', '') in lab.replace(' ', '') for k in SUBTOTAL_LABELS)):
            i += 1
            continue
        acc, j, done = (0.0, i + 1, False)
        while j < len(rows):
            r2, lab2, head2, v2 = rows[j]
            if head2.startswith('○') or any((k.replace(' ', '') in lab2.replace(' ', '') for k in SUBTOTAL_LABELS)):
                break
            if lab2 and v2:
                if done:
                    out.append(f"{r2}행 '{lab2}' ({v2 / 100000000.0:,.1f}억) - '{lab}' 소계에 안 잡힘 (계산에 미반영)")
                else:
                    acc += v2
                    if abs(acc - v) < max(1.0, abs(v) * 1e-06):
                        done = True
            j += 1
        i = j
    return out

def months_slice(ws, row, start_col):
    c0 = ci(start_col)
    return [num(ws.cell(row, c0 + i).value) for i in range(len(MONTHS))]
YM_HEADER = re.compile('^\\s*(\\d{2,4})\\s*년\\s*(\\d{1,2})\\s*월\\s*$')

def build_index(wb) -> tuple[dict, str]:
    ws = next((wb[n] for n in wb.sheetnames if 'INDEX' in n.upper()), None)
    if ws is None:
        return ({}, '')
    title = ''
    for r in range(1, min(ws.max_row, 5) + 1):
        for c in range(1, min(ws.max_column, 8) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and 'INDEX' in v.upper() and (len(v) > 8):
                title = v.strip()
                break
        if title:
            break
    headers = {}
    for r in range(1, ws.max_row + 1):
        cols = {}
        for c in range(1, ws.max_column + 1):
            m = YM_HEADER.match(str(ws.cell(r, c).value or ''))
            if m:
                y = int(m.group(1))
                cols[c] = f'{(y if y > 100 else 2000 + y):04d}-{int(m.group(2)):02d}'
        if cols:
            headers[r] = cols
    index = {}
    for r in range(1, ws.max_row + 1):
        labels = {str(ws.cell(r, c).value or '').strip() for c in range(1, 4)}
        if '환율' not in labels:
            continue
        hr = max((h for h in headers if h < r), default=None)
        if hr is None:
            continue
        fuel_row = next((rr for rr in range(r + 1, min(r + 4, ws.max_row + 1)) if '유가' in {str(ws.cell(rr, c).value or '').strip() for c in range(1, 4)}), None)
        for c, ym in headers[hr].items():
            fx = ws.cell(r, c).value
            fuel = ws.cell(fuel_row, c).value if fuel_row else None
            if isinstance(fx, (int, float)) and isinstance(fuel, (int, float)):
                index[ym] = {'fx': float(fx), 'fuel': float(fuel)}
    return (index, title)

def build(xlsx_path: Path) -> dict:
    global MONTHS
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    global ALLOC_SHEET, DIRECT_SHEET
    ALLOC_SHEET = next((s for s in wb.sheetnames if s.endswith(ALLOC_SUFFIX)), '')
    DIRECT_SHEET = next((s for s in wb.sheetnames if s.endswith(DIRECT_SUFFIX)), '')
    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    missing += [f'...{ALLOC_SUFFIX}'] if not ALLOC_SHEET else []
    missing += [f'...{DIRECT_SUFFIX}'] if not DIRECT_SHEET else []
    if missing:
        raise ValueError('필수 시트가 없습니다: ' + ', '.join(missing) + '\n   (시트 이름이 바뀌었다면 build_dataset.py 상단 목록을 맞춰주세요)')
    MONTHS = detect_months(wb)
    ref = wb['기준데이터']
    cask = wb['기종별 노선별 CASK']
    plan = wb['사업량']
    rev = wb['수입']
    alloc = wb[ALLOC_SHEET]
    anc = wb['부대수입']
    aircraft = {}
    for r in range(2, 20):
        code, name, seats = (ref.cell(r, 1).value, ref.cell(r, 2).value, ref.cell(r, 4).value)
        if not code or not name or (not isinstance(seats, (int, float))):
            continue
        if '공통' in str(name):
            continue
        aircraft[str(code).strip()] = {'code': str(code).strip(), 'name': str(name).strip(), 'seats': int(seats)}
    alloc_warn = resolve_alloc_rows(alloc, aircraft.keys())
    routes = {}
    for r in range(2, ref.max_row + 1):
        name = ref.cell(r, 6).value
        if not name:
            continue
        routes[str(name).strip()] = {'route': str(name).strip(), 'domestic_intl': ref.cell(r, 7).value, 'region': ref.cell(r, 8).value, 'origin': ref.cell(r, 9).value, 'distance_km': num(ref.cell(r, 10).value), 'block_time': num(ref.cell(r, 11).value)}
    name2code = {v['name']: k for k, v in aircraft.items()}
    cask_table = {}
    for r in range(5, cask.max_row + 1):
        ac_name, route = (cask.cell(r, 2).value, cask.cell(r, 3).value)
        if not ac_name or not route:
            continue
        code = name2code.get(str(ac_name).strip())
        if not code:
            continue
        key = f'{code}|{str(route).strip()}'
        if key in cask_table:
            continue
        items = {CASK_ITEM_NAMES[c]: num(cask.cell(r, ci(c)).value) for c in [CASK_FUEL_COL] + CASK_FX_COLS + CASK_KRW_COLS}
        cask_table[key] = {'aircraft': code, 'route': str(route).strip(), 'fuel': num(cask.cell(r, ci(CASK_FUEL_COL)).value), 'fx': sum((num(cask.cell(r, ci(c)).value) for c in CASK_FX_COLS)), 'krw': sum((num(cask.cell(r, ci(c)).value) for c in CASK_KRW_COLS)), 'items': items}

    def is_route_row(r):
        return bool(plan.cell(r, 5).value) and bool(plan.cell(r, 6).value) and (plan.cell(r, 2).value in ('국내선', '국제선'))
    route_rows = [r for r in range(5, 231) if is_route_row(r)]
    zeros = lambda: [0.0] * len(MONTHS)
    net = {k: zeros() for k in ['fc', 'bt', 'pax', 'ask', 'pax_rev', 'fc_lease', 'bt_lease']}
    by_type = {}
    for r in route_rows:
        region = plan.cell(r, 4).value
        if region == '부정기':
            continue
        code = str(plan.cell(r, 6).value).strip()
        fc = months_slice(plan, r, COL['사업량_편수'])
        bt = months_slice(plan, r, COL['사업량_BT'])
        pax = months_slice(plan, r, COL['사업량_수송석'])
        ask = months_slice(plan, r, COL['사업량_ASK'])
        prv = months_slice(rev, r, COL['수입_운송수입'])
        for i in range(len(MONTHS)):
            net['fc'][i] += fc[i]
            net['bt'][i] += bt[i]
            net['pax'][i] += pax[i]
            net['ask'][i] += ask[i]
            net['pax_rev'][i] += prv[i]
            if code not in LEASE_EXCLUDED_TYPES:
                net['fc_lease'][i] += fc[i]
                net['bt_lease'][i] += bt[i]
        t = by_type.setdefault(code, {'fc': zeros(), 'bt': zeros()})
        for i in range(len(MONTHS)):
            t['fc'][i] += fc[i]
            t['bt'][i] += bt[i]
    pools = {k: months_slice(alloc, row, COL['배부비_pool']) for k, row in POOL_ROW.items()}
    dep_pools = {k: months_slice(alloc, row, COL['배부비_pool']) for k, row in DEP_ROW.items()}
    alloc_warn += audit_alloc_sheet(alloc, len(MONTHS))
    head = find_section_row(anc, 'RASK', col='B') or 26
    rask = {}
    for r in range(head + 1, head + 40):
        region = anc.cell(r, 3).value
        lab = str(region).strip() if region else ''
        if not lab:
            if rask:
                break
            continue
        if lab in ('구분', '합계', '총합', '총계'):
            continue
        rask.setdefault(lab, months_slice(anc, r, COL['부대수입_rask']))
    npl = find_label_row(anc, '노선비연계', col='C')
    if npl is None:
        raise ValueError("'부대수입' 시트에서 '노선비연계' 행을 찾지 못했습니다.")
    anc_pool = months_slice(anc, npl, COL['부대수입_rask'])
    rr, cc = ESCALATION_FUEL_ROWCOL
    esc_fuel = num(wb[DIRECT_SHEET].cell(rr, ci(cc)).value)
    rr, cc = ESCALATION_OTHER_ROWCOL
    esc_other = num(wb[ALLOC_SHEET].cell(rr, ci(cc)).value)
    index, index_title = build_index(wb)
    plan_by_route = {}
    plan_detail: dict = {}
    for r in route_rows:
        name = str(plan.cell(r, 5).value).strip()
        code = str(plan.cell(r, 6).value).strip()
        region = plan.cell(r, 4).value
        seg = months_slice(plan, r, COL['사업량_편수'])
        e = plan_by_route.setdefault(name, {'total': [0.0] * len(MONTHS), 'by_type': {}})
        t = e['by_type'].setdefault(code, [0.0] * len(MONTHS))
        for i in range(len(MONTHS)):
            e['total'][i] += seg[i]
            t[i] += seg[i]
        if region == '부정기':
            continue
        d = plan_detail.setdefault(name, {}).setdefault(code, {k: [0.0] * len(MONTHS) for k in ('fc', 'bt', 'pax', 'ask', 'pax_rev')})
        for key, src, col in (('fc', plan, COL['사업량_편수']), ('bt', plan, COL['사업량_BT']), ('pax', plan, COL['사업량_수송석']), ('ask', plan, COL['사업량_ASK']), ('pax_rev', rev, COL['수입_운송수입'])):
            vals = months_slice(src, r, col)
            for i in range(len(MONTHS)):
                d[key][i] += vals[i]
    return {'meta': {'source': xlsx_path.name, 'built_at': datetime.now().isoformat(timespec='seconds'), 'months': MONTHS, 'escalation': {'fuel': esc_fuel, 'other': esc_other}, 'index_title': index_title, 'lease_excluded_types': sorted(LEASE_EXCLUDED_TYPES)}, 'index': index, 'aircraft': aircraft, 'routes': routes, 'plan_by_route': plan_by_route, 'plan_detail': plan_detail, 'cask': cask_table, 'network': net, 'network_by_type': by_type, 'pools': pools, 'depreciation_pools': dep_pools, 'ancillary_rask': rask, 'ancillary_pool': anc_pool, 'alloc_warnings': alloc_warn, 'alloc_rows': {'pool': POOL_ROW.copy(), 'dep': DEP_ROW.copy()}}

def main():
    if len(sys.argv) > 1:
        xlsx = Path(sys.argv[1])
        if not xlsx.exists():
            sys.exit(f'[ERROR] 파일을 찾을 수 없습니다: {xlsx}')
    else:
        xlsx = find_cost_file(BASE)
        if xlsx is None:
            sys.exit(f'[ERROR] 비용추정용 엑셀을 찾지 못했습니다.\n   {BASE} 안에 두거나, python build_dataset.py <파일경로> 로 지정하세요.')
    print(f'[원본] {xlsx.name}')
    others = [f for f in list_cost_files(BASE) if f != xlsx]
    if others:
        print('  [주의] 비용파일 후보가 여러 개입니다. 가장 최근 수정본을 골랐습니다.')
        for f in others:
            print(f'     - 안 쓴 파일 : {f.name} ({datetime.fromtimestamp(f.stat().st_mtime):%y.%m.%d %H:%M})')
        print('     → 구 시즌 파일은 폴더 밖으로 옮기시는 편이 안전합니다.')
    try:
        data = build(xlsx)
    except ValueError as e:
        sys.exit(f'[ERROR] {e}')
    print(f"     시즌 월 : {', '.join(data['meta']['months'])}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'[OK] {OUT}')
    print(f"     기종 {len(data['aircraft'])}종 / 노선 {len(data['routes'])}개 / CASK {len(data['cask'])}조합")
    rows = data['alloc_rows']
    print(f'     배부비 행(라벨로 탐색) : ' + ', '.join((f'{k}={v}' for k, v in rows['pool'].items())))
    print(f'     기종 감가 행 : ' + ', '.join((f'{k}={v}' for k, v in rows['dep'].items())))
    if data['alloc_warnings']:
        print()
        print('  [확인 필요] 배부비 시트에서 짚어볼 점')
        for w in data['alloc_warnings']:
            print(f'     ! {w}')
        print('     → 새로 생긴 비용 항목이면 build_dataset.py 의 POOL_LABELS 에 추가하고')
        print('       engine.py 에 배부 기준(편수/BT/운송수입/수송객)을 연결해야 반영됩니다.')
if __name__ == '__main__':
    main()
