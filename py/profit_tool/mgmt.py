# -*- coding: utf-8 -*-
from __future__ import annotations
import io
import pickle
import re
from datetime import datetime
from pathlib import Path
import openpyxl
from .paths import DATA_DIR, PASSWORD, TOOL
CACHE = TOOL / 'data' / 'mgmt_cache.pkl'
FILE_PAT = re.compile('경영기획.*노선별\\s*실적')
SHEET = 'AD-HOC(전체 실적)'
COL = dict(line=1, kind=2, area=3, route=4, ac=5, month=6, fc=7, seats=8, pax=9, lf=10, ar=11, rev=12, pax_rev=13, anc=14, cargo=15, cost=16, var=17, fix=20, rt_rev=23, rt_cost=24, rt_var=25, rt_fix=27)
HEAD_KEY = {'route': '왕복노선', 'ac': '기종', 'month': '월', 'fc': '운항편수', 'seats': '공급석', 'pax': '수송석', 'lf': 'lf', 'ar': 'ar', 'rev': '총수입', 'pax_rev': '운송수입', 'anc': '부대수입', 'cargo': '화물수입', 'cost': '총비용', 'var': '변동비', 'fix': '고정비', 'area': '대노선', 'kind': '정기', 'line': '국내/국제'}

def find_file() -> Path | None:
    if DATA_DIR is None:
        return None
    cands = [p for p in DATA_DIR.glob('*.xls*') if FILE_PAT.search(p.name) and (not p.name.startswith('~$'))]
    return max(cands, key=lambda p: p.stat().st_mtime) if cands else None

def _open(path: Path):
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except Exception:
        pass
    try:
        import msoffcrypto
    except ImportError:
        return None
    buf = io.BytesIO()
    try:
        with path.open('rb') as f:
            office = msoffcrypto.OfficeFile(f)
            office.load_key(password=PASSWORD or '')
            office.decrypt(buf)
        buf.seek(0)
        return openpyxl.load_workbook(buf, data_only=True)
    except Exception:
        return None

def _headers(rows) -> dict:
    for r in rows[:20]:
        txt = {i: str(c).strip().lower() for i, c in enumerate(r) if c}
        if any((v == '왕복노선' for v in txt.values())):
            col = {}
            for name, key in HEAD_KEY.items():
                col[name] = next((i for i, v in txt.items() if key in v), COL.get(name))
            for k in ('rt_rev', 'rt_cost', 'rt_var', 'rt_fix'):
                col[k] = COL[k]
            hits = {i: v for i, v in txt.items() if v.startswith('1왕복')}
            for k, key in (('rt_rev', '수입'), ('rt_cost', '비용'), ('rt_var', '변동비'), ('rt_fix', '고정비')):
                got = next((i for i, v in hits.items() if key in v), None)
                if got is not None:
                    col[k] = got
            return col
    return dict(COL)

def _year(name: str) -> int:
    m = re.search('(20\\d{2})년', name) or re.search('(\\d{2})년', name)
    if not m:
        return datetime.now().year
    y = int(m.group(1))
    return y if y > 100 else 2000 + y

def _parse(path: Path) -> dict:
    wb = _open(path)
    if wb is None or SHEET not in wb.sheetnames:
        return {}
    ws = wb[SHEET]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    col = _headers(rows)
    year = _year(path.name)
    by_ac, carry = ({}, {})
    for r in rows:
        for k in ('line', 'kind', 'area', 'route', 'ac'):
            i = col[k]
            if i is None or i >= len(r):
                continue
            if r[i] not in (None, ''):
                carry[k] = str(r[i]).strip()
            else:
                r[i] = carry.get(k)
        get = lambda k: r[col[k]] if col.get(k) is not None and col[k] < len(r) else None
        month, route, ac = (get('month'), get('route'), get('ac'))
        if not isinstance(month, (int, float)):
            continue
        if not route or not isinstance(get('fc'), (int, float)):
            continue
        if '기타' in str(route) or '기타' in str(get('area') or ''):
            continue
        if not (get('seats') or 0):
            continue
        key = (str(route).strip(), str(ac or '').strip(), year, int(month))
        by_ac[key] = {k: float(get(k)) if isinstance(get(k), (int, float)) else None for k in ('fc', 'seats', 'pax', 'lf', 'ar', 'rev', 'pax_rev', 'anc', 'cargo', 'cost', 'var', 'fix', 'rt_cost', 'rt_var', 'rt_fix')}
    return by_ac

class Mgmt:

    def __init__(self, path: Path | None=None):
        self.path = path or find_file()
        self.available = False
        self.label = ''
        self.months: list[tuple[int, int]] = []
        self._ac: dict = {}
        self._route: dict = {}
        if self.path is None or not self.path.exists():
            return
        st = self.path.stat()
        sig = f'{self.path.name}:{st.st_size}:{int(st.st_mtime)}'
        data = None
        try:
            cached = pickle.loads(CACHE.read_bytes())
            if cached.get('sig') == sig:
                data = cached['data']
        except Exception:
            pass
        if data is None:
            data = _parse(self.path)
            try:
                CACHE.parent.mkdir(parents=True, exist_ok=True)
                CACHE.unlink(missing_ok=True)
                CACHE.write_bytes(pickle.dumps({'sig': sig, 'data': data}))
            except OSError:
                pass
        self._ac = data
        for (route, ac, y, m), v in data.items():
            e = self._route.setdefault((route, y, m), {k: 0.0 for k in ('fc', 'seats', 'pax', 'rev', 'pax_rev', 'anc', 'cargo', 'cost', 'var', 'fix')})
            for k in e:
                e[k] += v.get(k) or 0.0
        self.months = sorted({(y, m) for _, _, y, m in data})
        self.available = bool(data)
        self.label = re.sub('\\.xls[xmb]?$', '', self.path.name, flags=re.I)

    def get(self, route: str, y: int, m: int) -> dict | None:
        return self._route.get((route, y, m))

    def by_aircraft(self, route: str, y: int, m: int) -> dict:
        return {ac: v for (r, ac, yy, mm), v in self._ac.items() if r == route and (yy, mm) == (y, m)}

    def covers(self, y: int, m: int) -> bool:
        return (y, m) in self.months

    @property
    def short(self) -> str:
        m = re.search('REV\\s*\\d+', self.label, re.I)
        return f"경영기획 {m.group(0).replace(' ', '')}" if m else '경영기획 노선별 실적'

    @property
    def note(self) -> str:
        if not self.available:
            return ''
        ms = self.months
        return f'{self.label} ({ms[0][0] % 100}.{ms[0][1]:02d}~{ms[-1][0] % 100}.{ms[-1][1]:02d} 확정)'
