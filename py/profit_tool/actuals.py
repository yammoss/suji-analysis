# -*- coding: utf-8 -*-
from __future__ import annotations
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import openpyxl
from .period import Period
BASE = Path(__file__).resolve().parent.parent
from .paths import DATA_DIR
DEFAULT_ACTUALS = DATA_DIR / '과거실적 DATA.xlsx'
C_YM, C_ROUTE, C_AC, C_FC, C_SEATS, C_PAX, C_PAXREV, C_CARGO = (1, 3, 4, 5, 6, 7, 11, 13)
HEADER_ALIASES = {'ym': ('년월', '연월', '년/월', '기준월', '해당월', 'ym', '년도월'), 'day': ('출발일', '출발일자', '운항일', '운항일자', '일자', 'date'), 'route': ('노선', '구간', 'route'), 'dow': ('요일', 'dow', 'weekday'), 'ac': ('a/c(소)', 'a/c', 'ac', '기종', '기종(소)', 'aircraft'), 'fc': ('운항편수', '편수', '운항', 'fc', 'seg'), 'seats': ('공급좌석', '공급석', '좌석', 'ask좌석', '공급'), 'pax': ('수송석', '수송객', '탑승객', '여객수', 'pax'), 'rev': ('운송수입', '여객수입', '수입', '매출'), 'cargo': ('화물수입', '화물'), 'lf': ('l/f', 'lf', '탑승률', '탑승율'), 'ar': ('a/r', 'ar', '운임', '단가')}
PLAN_HINTS = ('추정', '계획', '목표', '예상', '전망', 'plan', 'forecast', 'budget')
SUBTOTAL_NAMES = ('total', '합계', '소계', '계')
CARGO_SUBSTITUTE = {'A330-900': ['A330-200', 'A330-300'], 'A330-200': ['A330-300', 'A330-900'], 'A330-300': ['A330-200', 'A330-900'], 'B737-8': ['B737-800'], 'B737-800': ['B737-8'], 'B777-300ER': []}

@dataclass
class CargoRevenue:
    per_round_trip: float
    basis: str
    months_used: int = 0
    months_needed: int = 0

    @property
    def short(self) -> str:
        short_cover = self.months_needed and self.months_used < self.months_needed
        part = f' ※{self.months_used}/{self.months_needed}개월' if short_cover else ''
        return f'화물 {self.per_round_trip / 1000:,.0f}천원/왕복 ({self.basis}){part}'

@dataclass
class ActualRate:
    lf: float
    ar: float
    route: str
    months_used: list[tuple[int, int]]
    requested: str
    years_back: int
    months_needed: int = 0
    kind: str = 'actual'
    sheet: str = ''
    dows: str = ''
    dow_fallback: bool = False
    matches: list = field(default_factory=list)
    match_note: str = ''
    unmatched: int = 0
    day_fallback: bool = False
    pickup_label: str = ''

    @property
    def is_daily(self) -> bool:
        return self.kind == 'daily'

    @property
    def is_pickup(self) -> bool:
        return self.kind == 'pickup'

    @property
    def base_range(self) -> str:
        ds = [b for m in self.matches for b in m.base]
        if not ds:
            return ''
        a, b = (min(ds), max(ds))
        return f'{a:%y.%m.%d}' if a == b else f'{a:%y.%m.%d}~{b:%y.%m.%d}'

    @property
    def full_cover(self) -> bool:
        if self.is_daily:
            return self.unmatched == 0
        return self.months_needed == 0 or len(self.months_used) >= self.months_needed

    @property
    def is_plan(self) -> bool:
        return self.kind == 'plan'

    @property
    def basis_label(self) -> str:
        a, b = (self.months_used[0], self.months_used[-1])
        if a == b:
            return f'{a[0] % 100}.{a[1]:02d}월'
        return f'{a[0] % 100}.{a[1]:02d}~{b[0] % 100}.{b[1]:02d}월'

    @property
    def plan_year(self) -> str:
        ys = sorted({y for y, _ in self.months_used})
        if not ys:
            return ''
        if len(ys) == 1:
            return f'{ys[0] % 100}년'
        return f'{ys[0] % 100}~{ys[-1] % 100}년'

    @property
    def source_label(self) -> str:
        if self.is_plan:
            return f'{self.plan_year} 사업계획 목표실적'
        if self.is_daily:
            return f'{self.base_range} 일자 대응 실적'
        if self.is_pickup:
            return self.pickup_label
        tag = f'({self.dows})' if self.dows else ''
        if self.dow_fallback:
            tag = '(요일 실적 없어 전 요일)'
        return f'{self.basis_label} 실적{tag}'

    @property
    def short(self) -> str:
        lack = '개월만 계획 보유' if self.is_plan else '개월만 실적 보유'
        part = '' if self.full_cover else f' ※{len(self.months_used)}/{self.months_needed}{lack}'
        if self.is_daily and self.unmatched:
            part = f' ※운항일 {self.unmatched}일 대응 실적 없음'
        return f'L/F {self.lf:.1%} · A/R {self.ar:,.0f}원 ({self.source_label}){part}'

    @property
    def note(self) -> str:
        part = '' if self.full_cover else ' (일부 월만 보유)'
        return f'L/F·A/R = {self.source_label}{part}'

class Actuals:

    def __init__(self, path: Path | str=DEFAULT_ACTUALS):
        self.path = Path(path)
        self.available = self.path.exists()
        self._agg: dict[tuple[str, int, int], list[float]] = {}
        self._cargo: dict[tuple[str, str, int, int], list[float]] = {}
        self._agg_dow: dict[tuple, list[float]] = {}
        self._cargo_dow: dict[tuple, list[float]] = {}
        self._plan: dict[tuple[str, int, int], list[float]] = {}
        self._plan_cargo: dict[tuple[str, str, int, int], list[float]] = {}
        self._day: dict[str, dict[date, list[float]]] = {}
        self.day_range: tuple[date, date] | None = None
        self.plan_sheets: list[str] = []
        if not self.available:
            return
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        for i, ws in enumerate(wb.worksheets):
            plan = is_plan_sheet(ws.title)
            if i and (not plan):
                continue
            if plan:
                self.plan_sheets.append(ws.title)
                if self._read_matrix(ws):
                    continue
            self._read_sheet(ws, plan)
        wb.close()

    def _read_matrix(self, ws) -> bool:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        vals: dict[tuple[str, int, int], dict] = {}
        for hi, row in enumerate(rows):
            for rc, cell in enumerate(row):
                if not (isinstance(cell, str) and re.sub('\\s', '', cell) == '노선'):
                    continue
                months, j = ([], rc + 1)
                while j < len(row):
                    ym = _month_header(row[j])
                    if ym is None:
                        break
                    months.append((j, ym))
                    j += 1
                if not months:
                    continue
                kind = _block_kind(rows, hi, rc, months[-1][0])
                if not kind:
                    continue
                for r in rows[hi + 1:]:
                    route = r[rc] if rc < len(r) else None
                    if not isinstance(route, str) or not route.strip():
                        continue
                    for col, (y, m) in months:
                        v = r[col] if col < len(r) else None
                        v = _pct(v) if kind == 'lf' else _f(v)
                        if v > 0:
                            vals.setdefault((route.strip(), y, m), {})[kind] = v
        n = 0
        for key, d in vals.items():
            if d.get('lf') and d.get('ar'):
                a = self._plan.setdefault(key, [0.0] * 6)
                a[3] += d['lf']
                a[4] += d['ar']
                a[5] += 1.0
                n += 1
        return n > 0

    def _read_sheet(self, ws, plan: bool) -> None:
        col = None
        for row in ws.iter_rows(values_only=True):
            if col is None:
                col = _find_cols(row)
                continue
            ym, route = (_at(row, col['ym']), _at(row, col['route']))
            if ym is None or not route:
                continue
            ymd = _ym(ym)
            if ymd is None:
                continue
            y, m = ymd
            route = str(route).strip()
            if route.lower() in SUBTOTAL_NAMES:
                continue
            seats, pax = (_f(_at(row, col['seats'])), _f(_at(row, col['pax'])))
            rev, cargo = (_f(_at(row, col['rev'])), _f(_at(row, col['cargo'])))
            fc = _f(_at(row, col['fc']))
            ac = str(_at(row, col['ac']) or '').strip()
            dow = None if plan else _dow(_at(row, col['dow']))
            store_c = self._plan_cargo if plan else self._cargo
            if (cargo or fc) and ac and (ac != '-'):
                e = store_c.setdefault((route, ac, y, m), [0.0, 0.0])
                e[0] += cargo
                e[1] += fc
                if dow:
                    e = self._cargo_dow.setdefault((route, ac, y, m, dow), [0.0, 0.0])
                    e[0] += cargo
                    e[1] += fc
            if not plan:
                d = _at(row, col['day']) if col['day'] is not None else None
                if isinstance(d, (datetime, date)):
                    d = d.date() if isinstance(d, datetime) else d
                    lo, hi = self.day_range or (d, d)
                    self.day_range = (min(lo, d), max(hi, d))
                    if seats > 0:
                        a = self._day.setdefault(route, {}).setdefault(d, [0.0, 0.0, 0.0, 0.0])
                        a[0] += fc
                        a[1] += seats
                        a[2] += pax
                        a[3] += rev
                if seats <= 0:
                    continue
                keys = [(self._agg, (route, y, m))]
                if dow:
                    keys.append((self._agg_dow, (route, y, m, dow)))
                for store, key in keys:
                    a = store.setdefault(key, [0.0, 0.0, 0.0])
                    a[0] += seats
                    a[1] += pax
                    a[2] += rev
                continue
            lf, ar = (_pct(_at(row, col['lf'])), _f(_at(row, col['ar'])))
            if seats <= 0 and (not (lf and ar)):
                continue
            w = seats or fc or 1.0
            a = self._plan.setdefault((route, y, m), [0.0] * 6)
            a[0] += seats
            a[1] += pax
            a[2] += rev
            if lf and ar:
                a[3] += lf * w
                a[4] += ar * w
                a[5] += w

    @property
    def has_plan(self) -> bool:
        return bool(self._plan)

    @property
    def has_daily(self) -> bool:
        return bool(self._day)

    def for_dates_daily(self, route: str, period: Period, targets: list, calendar, years_back: int) -> 'ActualRate | None':
        from .daymatch import match_dates, summarize
        days = self._day.get(route)
        if not days or not targets:
            return None
        ms = match_dates(days, calendar, targets, years_back)
        used = [m for m in ms if m.base]
        seats = sum((m.seats for m in used))
        pax = sum((m.pax for m in used))
        rev = sum((m.rev for m in used))
        if seats <= 0 or pax <= 0:
            return None
        months = sorted({(b.year, b.month) for m in used for b in m.base})
        return ActualRate(lf=pax / seats, ar=rev / pax, route=route, months_used=months, requested=period.label, years_back=years_back, months_needed=len(period.months), kind='daily', matches=ms, match_note=summarize(ms, calendar), unmatched=len(ms) - len(used))

    @property
    def has_dow(self) -> bool:
        return bool(self._agg_dow)

    def _month_agg(self, route, y, m, dows, dows_by_month=None):
        eff = dows_by_month.get(m) if dows_by_month is not None else dows
        if not eff or not self._agg_dow:
            return self._agg.get((route, y, m))
        parts = [self._agg_dow.get((route, y, m, d)) for d in eff]
        parts = [x for x in parts if x]
        if not parts:
            return None
        return [sum((x[i] for x in parts)) for i in range(3)]

    def _month_cargo(self, route, ac, y, m, dows, dows_by_month=None):
        eff = dows_by_month.get(m) if dows_by_month is not None else dows
        if not eff or not self._cargo_dow:
            return self._cargo.get((route, ac, y, m))
        parts = [self._cargo_dow.get((route, ac, y, m, d)) for d in eff]
        parts = [x for x in parts if x]
        if not parts:
            return None
        return [sum((x[i] for x in parts)) for i in range(2)]

    def _plan_rate(self, route: str, period: Period) -> ActualRate | None:
        seats = pax = rev = lfw = arw = w = 0.0
        used = []
        for y, m in period.months:
            a = self._plan.get((route, y, m))
            if not a:
                continue
            seats += a[0]
            pax += a[1]
            rev += a[2]
            lfw += a[3]
            arw += a[4]
            w += a[5]
            used.append((y, m))
        if not used:
            return None
        if seats > 0 and pax > 0 and (rev > 0):
            lf, ar = (pax / seats, rev / pax)
        elif w > 0:
            lf, ar = (lfw / w, arw / w)
        else:
            return None
        return ActualRate(lf=lf, ar=ar, route=route, months_used=used, requested=period.label, years_back=0, months_needed=len(period.months), kind='plan', sheet=', '.join(self.plan_sheets))

    def cargo_revenue(self, route: str, ac_name: str, period: Period, years_back: int=1, max_back: int=3, dows: list | None=None, dows_by_month: dict | None=None) -> CargoRevenue:
        need = len(period.months)
        if self._plan_cargo:
            for key, label in [(ac_name, ac_name)] + [(a, f'{a} 대체') for a in CARGO_SUBSTITUTE.get(ac_name, [])]:
                cargo = fc = 0.0
                used = []
                for y, m in period.months:
                    e = self._plan_cargo.get((route, key, y, m))
                    if not e:
                        continue
                    cargo += e[0]
                    fc += e[1]
                    used.append((y, m))
                if used and fc > 0:
                    return CargoRevenue(cargo / fc * 2, f'{used[0][0] % 100}년 사업계획 목표 {label}', len(used), need)
        attempts = [(ac_name, ac_name)]
        for alt in CARGO_SUBSTITUTE.get(ac_name, []):
            attempts.append((alt, f'{alt} 대체'))
        for key, label in attempts:
            best = None
            for back in range(years_back, max_back + 1):
                src = period.shift_years(-back)
                cargo = fc = 0.0
                used = []
                for y, m in src.months:
                    e = self._month_cargo(route, key, y, m, dows, dows_by_month)
                    if not e:
                        continue
                    cargo += e[0]
                    fc += e[1]
                    used.append((y, m))
                if not used or fc <= 0:
                    continue
                a, b = (used[0], used[-1])
                span = f'{a[0] % 100}.{a[1]:02d}' if a == b else f'{a[0] % 100}.{a[1]:02d}~{b[0] % 100}.{b[1]:02d}'
                tag = _dow_tag(dows, dows_by_month, self._cargo_dow)
                res = CargoRevenue(cargo / fc * 2, f'{span}월 {label} 실적{tag}', len(used), need)
                if len(used) == need:
                    return res
                if best is None or best.months_used < len(used):
                    best = res
            if best:
                return best
        return CargoRevenue(0.0, '실적 없음', 0, need)

    def for_period(self, route: str, period: Period, years_back: int=1, max_back: int=3, use_plan: bool=True, dows: list | None=None, dows_by_month: dict | None=None) -> ActualRate | None:
        need = len(period.months)
        plan = self._plan_rate(route, period) if use_plan else None
        if plan and plan.full_cover:
            return plan
        best = None
        for back in range(years_back, max_back + 1):
            src = period.shift_years(-back)
            seats = pax = rev = 0.0
            used = []
            for y, m in src.months:
                a = self._month_agg(route, y, m, dows, dows_by_month)
                if not a:
                    continue
                seats += a[0]
                pax += a[1]
                rev += a[2]
                used.append((y, m))
            if not used or seats <= 0 or pax <= 0:
                continue
            rate = ActualRate(lf=pax / seats, ar=rev / pax, route=route, months_used=used, requested=period.label, years_back=back, months_needed=need, dows=_dow_tag(dows, dows_by_month, self._agg_dow, bare=True))
            if len(used) == need:
                return rate
            if best is None or len(used) > len(best.months_used):
                best = rate
        if plan and (best is None or len(plan.months_used) >= len(best.months_used)):
            return plan
        return best

def _dow_tag(dows, dows_by_month, has_dow_data, bare: bool=False) -> str:
    if not has_dow_data:
        return ''
    if dows_by_month is not None:
        used = {tuple(sorted(v)) if v else None for v in dows_by_month.values()}
        if len(used) > 1:
            txt = '월별로 다름'
        else:
            only = next(iter(used)) if used else None
            if only is None:
                return ''
            txt = '·'.join(only)
    elif dows:
        txt = '·'.join(dows)
    else:
        return ''
    return txt if bare else f'({txt})'

def _f(v):
    return float(v) if isinstance(v, (int, float)) else 0.0

def _pct(v):
    if isinstance(v, str):
        t = v.strip().rstrip('%').replace(',', '')
        try:
            v = float(t)
        except ValueError:
            return 0.0
    if not isinstance(v, (int, float)):
        return 0.0
    v = float(v)
    return v / 100 if v > 1.5 else v

def _at(row, idx):
    return None if idx is None or idx >= len(row) else row[idx]

def _ym(v):
    try:
        return (v.year, v.month)
    except AttributeError:
        pass
    import re as _re
    nums = _re.findall('\\d+', str(v))
    if len(nums) < 2:
        return None
    y, m = (int(nums[0]), int(nums[1]))
    if y < 100:
        y += 2000
    if not 1 <= m <= 12 or not 2000 <= y <= 2099:
        return None
    return (y, m)
_DOW_NAMES = ('월', '화', '수', '목', '금', '토', '일')

def _dow(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and 1 <= int(v) <= 7:
        return _DOW_NAMES[int(v) - 1]
    t = str(v).strip()
    if t[:1] in _DOW_NAMES:
        return t[:1]
    eng = {'MON': '월', 'TUE': '화', 'WED': '수', 'THU': '목', 'FRI': '금', 'SAT': '토', 'SUN': '일'}
    return eng.get(t[:3].upper())

def _month_header(v):
    if v is None or isinstance(v, (int, float)):
        return None
    if not hasattr(v, 'year') and '월' not in str(v) and (not re.search('\\d[.\\-/]\\d', str(v))):
        return None
    return _ym(v)

def _block_kind(rows, hi: int, first_col: int, last_col: int):
    lo = max(0, first_col - 4)
    for r in range(hi, max(-1, hi - 5), -1):
        for c in range(lo, min(len(rows[r]), last_col + 1)):
            t = str(rows[r][c] or '').upper().replace(' ', '')
            if 'A/R' in t or '운임' in t:
                return 'ar'
            if 'L/F' in t or '탑승률' in t or '탑승율' in t:
                return 'lf'
    return None

def is_plan_sheet(title: str) -> bool:
    t = str(title).lower()
    return any((h in t for h in PLAN_HINTS))

def _find_cols(header) -> dict:
    found = {}
    for i, cell in enumerate(header or ()):
        if cell is None:
            continue
        key = re.sub('[\\s()]', '', str(cell)).lower()
        for name, names in HEADER_ALIASES.items():
            if name in found:
                continue
            if key in tuple((re.sub('[\\s()]', '', n).lower() for n in names)):
                found[name] = i
                break
    if 'ym' not in found or 'route' not in found:
        found = {'ym': C_YM - 1, 'route': C_ROUTE - 1, 'ac': C_AC - 1, 'fc': C_FC - 1, 'seats': C_SEATS - 1, 'pax': C_PAX - 1, 'rev': C_PAXREV - 1, 'cargo': C_CARGO - 1}
    for name in HEADER_ALIASES:
        found.setdefault(name, None)
    return found
