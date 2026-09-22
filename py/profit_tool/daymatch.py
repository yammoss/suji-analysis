# -*- coding: utf-8 -*-
from __future__ import annotations
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
BASE = Path(__file__).resolve().parent.parent
from .paths import DATA_DIR
DEFAULT_CALENDAR = DATA_DIR / '공휴일 DATA.xlsx'
DAY_NAMES = '월화수목금토일'
BIG = ('설날', '추석')
WINDOW = 2
WEEKS = 364

def fd(d: date) -> str:
    return f'{d:%y.%m.%d}({DAY_NAMES[d.weekday()]})'

class HolidayCalendar:

    def __init__(self, path: Path | str=DEFAULT_CALENDAR, country: str='한국'):
        self.path = Path(path)
        self.country = country
        self.hol: dict[date, str] = {}
        self.estimated: set[date] = set()
        self.years: set[int] = set()
        self.available = self.path.exists()
        if self.available:
            self._load()
        self._index()

    def _load(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        if '공휴일목록' not in wb.sheetnames:
            self.available = False
            wb.close()
            return
        rows = wb['공휴일목록'].iter_rows(values_only=True)
        head = [str(h or '') for h in next(rows)]
        col = lambda name: next((i for i, h in enumerate(head) if h.startswith(name)))
        c_cty, c_day, c_kind, c_ev, c_st = (col('국가'), col('날짜'), col('종류'), col('이벤트'), col('상태'))
        for r in rows:
            d = r[c_day]
            if r[c_cty] != self.country or not isinstance(d, (datetime, date)):
                continue
            d = d.date() if isinstance(d, datetime) else d
            self.years.add(d.year)
            if r[c_kind] != '휴일':
                continue
            self.hol[d] = str(r[c_ev])
            if r[c_st] == '추정':
                self.estimated.add(d)
        wb.close()
        self.available = bool(self.hol)

    def is_off(self, d: date) -> bool:
        return d in self.hol or d.weekday() >= 5

    def covers(self, d: date) -> bool:
        return d.year in self.years

    def _index(self):
        self.blocks: list[tuple[str, list[date]]] = []
        self.by_date: dict[date, tuple] = {}
        seen = set()
        for h in sorted(self.hol):
            if h in seen:
                continue
            s = e = h
            while self.is_off(s - timedelta(1)):
                s -= timedelta(1)
            while self.is_off(e + timedelta(1)):
                e += timedelta(1)
            span = [s + timedelta(i) for i in range((e - s).days + 1)]
            seen.update(span)
            evs = [self.hol[x] for x in span if x in self.hol]
            ev = next((x for x in evs if x in BIG), evs[0])
            self.blocks.append((ev, span))
            for i, d in enumerate(span):
                if ev in BIG or d in self.hol:
                    self.by_date[d] = (ev, span, i, '연휴')
            if ev in BIG:
                for k in range(1, WINDOW + 1):
                    self.by_date.setdefault(span[0] - timedelta(k), (ev, span, -k, '직전'))
                    self.by_date.setdefault(span[-1] + timedelta(k), (ev, span, k, '직후'))

    def event_block(self, ev: str, near: date, within: int=60):
        cands = [sp for e, sp in self.blocks if e == ev and abs((sp[0] - near).days) <= within]
        return min(cands, key=lambda sp: abs((sp[0] - near).days)) if cands else None

@dataclass
class DayMatch:
    target: date
    label: str
    base: list = field(default_factory=list)
    why: str = ''
    fc: float = 0.0
    seats: float = 0.0
    pax: float = 0.0
    rev: float = 0.0
    exception: bool = False
    event: str = ''
    replaced: bool = False
    estimated: bool = False

def season_offset(day_range, first: date, last: date, max_back: int=3) -> int | None:
    if not day_range:
        return None
    lo, hi = day_range
    for k in range(1, max_back + 1):
        if first - timedelta(WEEKS * k + 14) >= lo and last - timedelta(WEEKS * k) + timedelta(14) <= hi:
            return k
    for k in range(1, max_back + 1):
        if last - timedelta(WEEKS * k) >= lo and first - timedelta(WEEKS * k) <= hi:
            return k
    return None

def match_dates(days: dict, cal: HolidayCalendar, targets: list, k: int) -> list[DayMatch]:
    shift = timedelta(WEEKS * k)

    def usable(b):
        return b in days and b not in cal.by_date
    out = []
    for t in sorted(set(targets)):
        near = t - shift
        est = t in cal.estimated or near in cal.estimated or (not cal.covers(t)) or (not cal.covers(near))
        info = cal.by_date.get(t)
        m = None
        if info:
            ev, span, pos, kind = info
            bspan = cal.event_block(ev, near)
            if bspan:
                if kind == '연휴':
                    label = f'{ev} 연휴 {pos + 1}/{len(span)}일째' if ev in BIG else ev
                    if len(span) == len(bspan):
                        b = [bspan[pos]]
                    elif len(span) == 1:
                        b = [x for x in bspan if cal.hol.get(x) == ev] or [bspan[0]]
                    elif len(bspan) == 1 or pos == 0:
                        b = [bspan[0]]
                    elif pos == len(span) - 1:
                        b = [bspan[-1]]
                    else:
                        b = bspan[1:-1]
                    why = f'{ev} 연휴 대응 ({fd(bspan[0])}~{fd(bspan[-1])})' if ev in BIG else f"{ev} 대응 ({' · '.join((fd(x) for x in b))})"
                else:
                    label = f'{ev} 연휴 {kind} D{pos:+d}'
                    b = [bspan[0] + timedelta(pos)] if pos < 0 else [bspan[-1] + timedelta(pos)]
                    why = f'{ev} 연휴 {kind} 대응'
                b = [x for x in b if x in days]
                if b:
                    m = DayMatch(t, label, b, why, exception=True, event=ev, estimated=est)
        if m is None:
            label = '주말' if t.weekday() >= 5 else '평일'
            if info and info[3] == '연휴':
                label = info[0]
            if usable(near):
                m = DayMatch(t, label, [near], f'{k}년 전 같은 요일', estimated=est)
            else:
                for w in (7, 14, 21):
                    alt = [x for x in (near - timedelta(w), near + timedelta(w)) if usable(x)]
                    if alt:
                        cause = f'{fd(near)} {cal.by_date[near][0]} {cal.by_date[near][3]}' if near in cal.by_date else f'{fd(near)} 실적 없음'
                        m = DayMatch(t, label, alt, f'{cause} → 앞뒤 {w}일 같은 요일', exception=True, replaced=near in cal.by_date, estimated=est)
                        break
                else:
                    m = DayMatch(t, label, [], '대응 실적 없음', exception=True, estimated=est)
        if m.base:
            n = len(m.base)
            m.fc, m.seats, m.pax, m.rev = (sum((days[x][i] for x in m.base)) / n for i in range(4))
        out.append(m)
    return out

def summarize(ms: list[DayMatch], cal: HolidayCalendar) -> str:
    parts = []
    for ev in dict.fromkeys((m.event for m in ms if m.event)):
        mine = [m for m in ms if m.event == ev]
        if ev in BIG:
            bs = sorted({b for m in mine for b in m.base})
            blk = cal.event_block(ev, bs[0], within=10) if bs else None
            where = f'{blk[0]:%y.%m.%d}~{blk[-1]:%m.%d} 연휴·앞뒤' if blk else ''
            parts.append(f'{ev} 대응 {len(mine)}일 ({where})')
        else:
            bs = sorted({b for m in mine for b in m.base})
            parts.append(f"{ev} 대응 {len(mine)}일 ({' · '.join((f'{b:%y.%m.%d}' for b in bs))})")
    rep = sum((1 for m in ms if m.replaced))
    if rep:
        parts.append(f'대응일이 연휴라 앞뒤 주로 대체 {rep}일')
    gap = sum((1 for m in ms if m.base and m.exception and (not m.replaced) and (not m.event)))
    if gap:
        parts.append(f'대응일 실적 없어 앞뒤 주로 대체 {gap}일')
    miss = sum((1 for m in ms if not m.base))
    if miss:
        parts.append(f'대응 실적 없음 {miss}일')
    return ' / '.join(parts)
