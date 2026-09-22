# -*- coding: utf-8 -*-
from __future__ import annotations
import bisect
import pickle
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from .paths import DATA_DIR
BASE = Path(__file__).resolve().parent.parent
CACHE = BASE / 'data' / 'booking_cache.pkl'
BUCKETS = ['M-5+', 'M-4', 'M-3', 'M-2', 'M-1', 'M-0']

def find_file() -> Path | None:
    cands = [p for p in DATA_DIR.glob('*.xlsx') if '발매' in p.name and (not p.name.startswith('~$'))]
    return max(cands, key=lambda p: p.stat().st_mtime) if cands else None

def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0

def bucket(sale: date, dy: int, dm: int) -> str:
    n = dy * 12 + dm - (sale.year * 12 + sale.month)
    return 'M-5+' if n >= 5 else f'M-{max(n, 0)}'

@dataclass
class Month:
    seats: float = 0.0
    dates: list = field(default_factory=list)
    cum: list = field(default_factory=list)
    rows: list = field(default_factory=list)

    def to_date(self, cutoff: date):
        i = bisect.bisect_right(self.dates, cutoff)
        return self.cum[i - 1] if i else (0.0, 0.0, 0.0, 0.0)

    @property
    def final(self):
        return self.cum[-1] if self.cum else (0.0, 0.0, 0.0, 0.0)

class Bookings:

    def __init__(self, path: Path | None=None, quiet: bool=False):
        self.path = path or find_file()
        self.available = bool(self.path and self.path.exists())
        self.months: dict[tuple, Month] = {}
        self.line: dict[str, str] = {}
        self.cutoff: date | None = None
        if self.available:
            self._load(quiet)

    def _load(self, quiet):
        st = self.path.stat()
        sig = (self.path.name, st.st_size, int(st.st_mtime))
        if CACHE.exists():
            try:
                with open(CACHE, 'rb') as f:
                    got = pickle.load(f)
                if got.get('sig') == sig:
                    self.months, self.line, self.cutoff = (got['months'], got['line'], got['cutoff'])
                    return
            except Exception:
                pass
        if not quiet:
            import sys
            print(f'   발매 현황 파일을 처음 읽는 중입니다 (1분 정도, 다음부터는 바로): {self.path.name}', file=sys.stderr)
        self._read_raw()
        CACHE.parent.mkdir(parents=True, exist_ok=True)
        CACHE.unlink(missing_ok=True)
        with open(CACHE, 'wb') as f:
            pickle.dump(dict(sig=sig, months=self.months, line=self.line, cutoff=self.cutoff), f)

    def _read_raw(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        ws = wb['RAW'] if 'RAW' in wb.sheetnames else wb.worksheets[-1]
        it = ws.iter_rows(values_only=True)
        head = [str(h or '').strip() for h in next(it)]
        c = {h: i for i, h in enumerate(head)}
        need = ('발매일자', '노선', '발매건수', '발매수입', '공급석', '출발연도', '출발월')
        miss = [k for k in need if k not in c]
        if miss:
            raise ValueError(f'발매 현황 RAW 시트에 {miss} 열이 없습니다')
        daily: dict[tuple, dict] = {}
        seats: dict[tuple, float] = {}
        last = None
        for r in it:
            d = r[c['발매일자']]
            route = r[c['노선']]
            if not route or not isinstance(d, (datetime, date)):
                continue
            d = d.date() if isinstance(d, datetime) else d
            try:
                key = (str(route).strip(), int(r[c['출발연도']]), int(r[c['출발월']]))
            except (TypeError, ValueError):
                continue
            last = d if last is None or d > last else last
            pax = _num(r[c['발매건수']])
            rev = _num(r[c['발매수입']])
            fuel = _num(r[c['유류수입']]) if '유류수입' in c else 0.0
            grp = pax if 'INDV/GRP' in c and str(r[c['INDV/GRP']]).lower().startswith('g') else 0.0
            a = daily.setdefault(key, {}).setdefault(d, [0.0, 0.0, 0.0, 0.0])
            a[0] += pax
            a[1] += rev
            a[2] += fuel
            a[3] += grp
            seats[key] = _num(r[c['공급석']]) or seats.get(key, 0.0)
            if '대노선' in c and r[c['대노선']]:
                self.line[key[0]] = str(r[c['대노선']]).strip()
        wb.close()
        self.cutoff = last
        for key, days in daily.items():
            m = Month(seats=seats.get(key, 0.0))
            tot = [0.0, 0.0, 0.0, 0.0]
            for d in sorted(days):
                v = days[d]
                m.rows.append((d, *v))
                tot = [tot[i] + v[i] for i in range(4)]
                m.dates.append(d)
                m.cum.append(tuple(tot))
            self.months[key] = m

    def complete(self, key) -> bool:
        y, m = (key[1], key[2])
        nxt = date(y + (m == 12), m % 12 + 1, 1)
        return self.cutoff is not None and nxt <= self.cutoff

    def by_bucket(self, key):
        out = {b: [0.0, 0.0, 0.0] for b in BUCKETS}
        m = self.months.get(key)
        if not m:
            return out
        for d, pax, rev, fuel, _ in m.rows:
            b = out[bucket(d, key[1], key[2])]
            b[0] += pax
            b[1] += rev
            b[2] += fuel
        return out

@dataclass
class Pickup:
    route: str
    ym: tuple
    cutoff: date
    seats: float
    booked_pax: float
    booked_rev: float
    ly_same_pax: float
    ly_same_rev: float
    ly_final_pax: float
    ly_final_rev: float
    ly_seats: float
    conv_pax: float
    conv_rev: float
    proj_pax: float
    proj_rev: float

    @property
    def lf(self):
        return self.proj_pax / self.seats if self.seats else None

    @property
    def ar(self):
        return self.proj_rev / self.proj_pax if self.proj_pax else None

    @property
    def booked_lf(self):
        return self.booked_pax / self.seats if self.seats else None

    @property
    def ly_same_lf(self):
        return self.ly_same_pax / self.ly_seats if self.ly_seats else None

    @property
    def pace(self):
        if not self.ly_same_lf:
            return None
        return self.booked_lf / self.ly_same_lf - 1

    @property
    def ahead(self) -> int:
        return self.ym[0] * 12 + self.ym[1] - (self.cutoff.year * 12 + self.cutoff.month)

def pickup_rate(bk: Bookings, route: str, period, flown: dict | None=None):
    from .actuals import ActualRate
    if not bk or not bk.available or len(period.months) != 1:
        return None
    y, m = period.months[0]
    p = pickup(bk, route, y, m, flown=flown)
    if not p or not p.lf or (not p.ar) or (not 0 <= p.ahead <= MAX_AHEAD):
        return None
    label = f'{p.cutoff:%y.%m.%d} 발매 기준 예측(발매율 {p.booked_lf:.0%}, 전년 동시점 {p.ly_same_lf:.0%})'
    return ActualRate(lf=p.lf, ar=p.ar, route=route, months_used=[(y, m)], requested=period.label, years_back=0, months_needed=1, kind='pickup', pickup_label=label)
MAX_AHEAD = 3

def pickup(bk: Bookings, route: str, y: int, m: int, cutoff: date | None=None, flown: dict | None=None) -> Pickup | None:
    cutoff = cutoff or bk.cutoff
    cur = bk.months.get((route, y, m))
    ly = bk.months.get((route, y - 1, m))
    if not cur or not ly or (not bk.complete((route, y - 1, m))) or (not cur.seats) or (not ly.seats):
        return None
    ly_cut = cutoff - timedelta(days=364)
    b_pax, b_rev, _, _ = cur.to_date(cutoff)
    s_pax, s_rev, _, _ = ly.to_date(ly_cut)
    f_pax, f_rev, _, _ = ly.final
    if f_pax <= 0:
        return None
    b, s, f = (b_pax / cur.seats, s_pax / ly.seats, f_pax / ly.seats)
    fill = max((f - s) / (1 - s), 0.0) if s < 1 else 0.0
    proj_pax = min(b + (1 - b) * fill, 1.02) * cur.seats if b < 1 else b_pax
    ly_rest_ar = (f_rev - s_rev) / (f_pax - s_pax) if f_pax > s_pax else f_rev / f_pax
    idx = b_rev / b_pax / (s_rev / s_pax) if b_pax > 0 and s_pax > 0 and (s_rev > 0) else 1.0
    proj_rev = b_rev + max(proj_pax - b_pax, 0.0) * ly_rest_ar * idx
    cp = cr = 1.0
    if flown and flown.get((route, y - 1, m)):
        fp, fr = flown[route, y - 1, m]
        cp = fp / f_pax if f_pax else 1.0
        cr = fr / f_rev if f_rev else 1.0
        cp, cr = (min(max(cp, 0.8), 1.05), min(max(cr, 0.8), 1.05))
    return Pickup(route, (y, m), cutoff, cur.seats, b_pax, b_rev, s_pax, s_rev, f_pax, f_rev, ly.seats, cp, cr, proj_pax * cp, proj_rev * cr)
