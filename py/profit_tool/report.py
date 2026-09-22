# -*- coding: utf-8 -*-
from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .engine import LegResult, Scenario
FONT_NAME = '대명체 굵게'
LINE_RGB = 'BFBFBF'
THIN = Side(style='thin', color=LINE_RGB)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOPLINE = Border(top=THIN, left=THIN, right=THIN, bottom=THIN)
HBAR = Border(top=THIN, bottom=THIN)
VBAR_L = Border(left=THIN, top=THIN, bottom=THIN)
VBAR_LR = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_TITLE = Font(name=FONT_NAME, size=20, bold=True)
F_HEAD = Font(name=FONT_NAME, size=10, bold=True)
F_BODY = Font(name=FONT_NAME, size=10)
F_SECT = Font(name=FONT_NAME, size=11, bold=True)
F_TEXT = Font(name=FONT_NAME, size=10)
F_INPUT = Font(name=FONT_NAME, size=10, color='0000C0')
F_NOTE = Font(name=FONT_NAME, size=9, italic=True)
F_SUM = Font(name=FONT_NAME, size=10, bold=True)
HEAD_FILL = PatternFill('solid', fgColor='666058')
F_HEAD_W = Font(name=FONT_NAME, size=10, bold=True, color='FFFFFF')
TITLE_FILL = PatternFill('solid', fgColor='C4B796')
F_TITLE_W = Font(name=FONT_NAME, size=20, bold=True, color='FFFFFF')
BANNER_FILL = HEAD_FILL
F_BANNER_W = Font(name=FONT_NAME, size=12, bold=True, color='FFFFFF')
WARN_FILL = PatternFill('solid', fgColor='FFF2CC')
NA_FILL = PatternFill('solid', fgColor='F8CBAD')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center')
ROW_H = 20.1
HEAD_H = 24.0
COLUMNS = [('기간', '', 15), ('노선', '', 17), ('편명', '', 11), ('기종', '', 8), ('CFG', '', 6), ('운항횟수', '(왕복)', 9), ('1왕복수입 (천원)', 'L/F(%)', 8), ('', 'A/R(원)', 10), ('', '여객', 13), ('', '부대', 11), ('', '화물', 11), ('', '총수입', 13), ('1왕복수지 (천원)', '변동비', 13), ('', '총비용', 13), ('', '한계이익', 13), ('', '영업이익', 13), ('', '한계이익률', 10), ('', '영업이익률', 10), ('총수지', '총수입', 11), ('', '총비용', 11), ('', '한계이익', 11), ('', '영업이익', 11), ('비고', '', 46)]
C_PERIOD, C_ROUTE, C_FLT, C_AC, C_CFG, C_RT = (2, 3, 4, 5, 6, 7)
C_LF, C_AR, C_PAX, C_ANC, C_CARGO, C_TOTREV = (8, 9, 10, 11, 12, 13)
C_VAR, C_TOTCOST, C_CM, C_OP, C_CMR, C_OPR = (14, 15, 16, 17, 18, 19)
C_PR, C_PC, C_PCM, C_POP = (20, 21, 22, 23)
C_NOTE = 24
LAST_COL = C_NOTE + 1
MONEY = '#,##0'
MONEY_K = '#,##0,;-#,##0,;'
MONEY_E = '#,##0"억";-#,##0"억";'
PL_K = '[Red]+#,##0,;[Blue]-#,##0,;[Black]"-"'
PL_E = '[Red]+#,##0"억";[Blue]-#,##0"억";[Black]"-"'
PCT = '[Red]"▲ +"0%;[Blue]"▼ -"0%;[Black]"-"'
FX_UNIT = '#,##0"원"'
FUEL_UNIT = '#,##0"c"'
EOK = 100000000
VLINE_COLS = {C_PERIOD, C_ROUTE, C_FLT, C_AC, C_CFG, C_RT, C_LF, C_VAR, C_PR, C_NOTE}
TOTAL_FILL = PatternFill('solid', fgColor='E6CBB8')

def cell_border(col: int) -> Border:
    if col == C_NOTE:
        return VBAR_LR
    return VBAR_L if col in VLINE_COLS else HBAR
PERIOD_PAIRS = [(C_PR, C_TOTREV), (C_PC, C_TOTCOST), (C_PCM, C_CM), (C_POP, C_OP)]
SUMMARY_ROWS = [('총수입', C_PR), ('총비용', C_PC), ('한계이익', C_PCM), ('영업이익', C_POP)]

def L(col: int) -> str:
    return get_column_letter(col)

def _text(ws, row, col, value, font=None, align=LEFT):
    c = ws.cell(row, col, value)
    c.font = font or F_TEXT
    c.alignment = align
    return c

def _write_header(ws, row: int) -> int:
    ws.row_dimensions[row].height = HEAD_H
    ws.row_dimensions[row + 1].height = HEAD_H
    spans, i = ([], 0)
    while i < len(COLUMNS):
        h1 = COLUMNS[i][0]
        j = i + 1
        if h1:
            while j < len(COLUMNS) and COLUMNS[j][0] == '':
                j += 1
        spans.append((i, j, h1))
        i = j
    for i, (h1, h2, _) in enumerate(COLUMNS, start=2):
        for r, v in ((row, None), (row + 1, h2 or None)):
            c = ws.cell(r, i, v)
            c.font = F_HEAD_W
            c.fill = HEAD_FILL
            c.alignment = CENTER
            c.border = cell_border(i)
    for a, b, h1 in spans:
        col = a + 2
        if b - a == 1 and (not COLUMNS[a][1]):
            ws.cell(row, col, h1).font = F_HEAD_W
            ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)
        else:
            ws.cell(row, col, h1).font = F_HEAD_W
            if b - a > 1:
                ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=b + 1)
        ws.cell(row, col).alignment = CENTER
    return row + 2

def _write_row(ws, row: int, r: LegResult):
    ws.row_dimensions[row].height = ROW_H
    n = str(row)
    P, X = (f'{L(C_PAX)}{n}', f'({L(C_CFG)}{n}*{L(C_LF)}{n}*2)')
    g = f'IF(OR({L(C_LF)}{n}="",{L(C_AR)}{n}=""),0,'
    vals = {C_PERIOD: r.period, C_ROUTE: r.leg.route, C_FLT: r.leg.flight_no, C_AC: r.leg.aircraft, C_CFG: r.seats, C_RT: r.leg.round_trips, C_LF: r.leg.lf, C_AR: r.leg.ar, C_CARGO: r.cargo_revenue if r.cargo_revenue is not None else 0, C_PAX: f'={g}{L(C_CFG)}{n}*{L(C_LF)}{n}*{L(C_AR)}{n}*2)', C_ANC: f'={g}{r.anc_const:.6f}+{r.anc_coef_rev:.12f}*{P})', C_TOTREV: f'={g}SUM({L(C_PAX)}{n}:{L(C_CARGO)}{n}))', C_VAR: f'={r.var_const:.6f}+{r.var_coef_rev:.12f}*{P}+{r.var_coef_pax:.6f}*{X}', C_TOTCOST: f'={L(C_VAR)}{n}+{r.fix_const:.6f}+{r.fix_coef_rev:.12f}*{P}', C_CM: f'={g}{L(C_TOTREV)}{n}-{L(C_VAR)}{n})', C_OP: f'={g}{L(C_TOTREV)}{n}-{L(C_TOTCOST)}{n})', C_CMR: f'=IFERROR({L(C_CM)}{n}/{L(C_TOTREV)}{n},"")', C_OPR: f'=IFERROR({L(C_OP)}{n}/{L(C_TOTREV)}{n},"")', C_NOTE: ' / '.join(r.notes)}
    for pcol, base in PERIOD_PAIRS:
        vals[pcol] = f'={g}{L(base)}{n}*{L(C_RT)}{n}/{EOK})'
    if not r.cask.available:
        for c in (C_VAR, C_TOTCOST, C_CM, C_OP, C_CMR, C_OPR, C_PC, C_PCM, C_POP):
            vals[c] = None
    for i in range(2, LAST_COL):
        v = vals.get(i)
        c = ws.cell(row, i, v if v not in ('', None) else None)
        c.border = cell_border(i)
        c.font = F_INPUT if i in (C_CFG, C_RT, C_LF, C_AR) else F_BODY
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True) if i == C_NOTE else CENTER
        if i == C_LF:
            c.number_format = '0.0%'
        elif i in (C_CMR, C_OPR):
            c.number_format = PCT
        elif i == C_AR:
            c.number_format = MONEY
        elif i in (C_CM, C_OP):
            c.number_format = PL_K
        elif i in (C_PCM, C_POP):
            c.number_format = PL_E
        elif i in (C_PR, C_PC):
            c.number_format = MONEY_E
        elif C_PAX <= i < C_NOTE:
            c.number_format = MONEY_K
    fill = None if r.cask.level == 'EXACT' else WARN_FILL if r.cask.available else NA_FILL
    if fill:
        for i in range(2, LAST_COL):
            ws.cell(row, i).fill = fill

def _sheet_ref(name: str) -> str:
    return "'" + name.replace("'", "''") + "'!"

def _rng(col, a, b, sheet: str | None=None):
    r = f'${L(col)}${a}:${L(col)}${b}'
    return _sheet_ref(sheet) + r if sheet else r
AVG_MONEY_COLS = (C_PAX, C_ANC, C_CARGO, C_TOTREV, C_VAR, C_TOTCOST, C_CM, C_OP)

def _agg_formulas(row: int, first: int, last: int, sheet: str | None=None, blank: str='""') -> dict:
    R = lambda col: _rng(col, first, last, sheet)
    rt, cfg, lf = (R(C_RT), R(C_CFG), R(C_LF))
    has = f'({lf}>0)'
    den = f'SUMPRODUCT({has}*{rt})'
    pax = f'SUMPRODUCT({has}*{cfg}*{lf}*{rt}*2)'
    seat = f'SUMPRODUCT({has}*{cfg}*{rt}*2)'
    f = {C_RT: f'=SUM({rt})', C_LF: f'=IFERROR({pax}/{seat},{blank})', C_AR: f'=IFERROR(SUMPRODUCT({has}*{R(C_PAX)}*{rt})/{pax},{blank})'}
    for col in AVG_MONEY_COLS:
        if col in (C_VAR, C_TOTCOST):
            f[col] = f'=IFERROR(SUMPRODUCT({has}*{R(col)}*{rt})/{den},IFERROR(SUMPRODUCT({R(col)}*{rt})/SUM({rt}),{blank}))'
        else:
            f[col] = f'=IFERROR(SUMPRODUCT({has}*{R(col)}*{rt})/{den},{blank})'
    for col, base in ((C_CMR, C_CM), (C_OPR, C_OP)):
        f[col] = f'=IFERROR({L(base)}{row}/{L(C_TOTREV)}{row},"")'
    for _, col in SUMMARY_ROWS:
        f[col] = f'=SUM({R(col)})'
    return f

def _write_total(ws, row: int, first: int, last: int):
    ws.row_dimensions[row].height = ROW_H
    for i in range(2, LAST_COL):
        c = ws.cell(row, i)
        c.border = cell_border(i)
        c.fill = TOTAL_FILL
        c.font = F_SUM
        c.alignment = CENTER
    ws.cell(row, C_PERIOD, '합계')
    for col, formula in _agg_formulas(row, first, last).items():
        c = ws.cell(row, col, formula)
        if col == C_RT or col == C_AR:
            c.number_format = MONEY
        elif col == C_LF:
            c.number_format = '0.0%'
        elif col in (C_CMR, C_OPR):
            c.number_format = PCT
        elif col in (C_CM, C_OP):
            c.number_format = PL_K
        elif col in (C_PCM, C_POP):
            c.number_format = PL_E
        elif col in (C_PR, C_PC):
            c.number_format = MONEY_E
        else:
            c.number_format = MONEY_K

def _write_link_row(ws, row: int, r: LegResult, sheet: str, first: int, last: int):
    ws.row_dimensions[row].height = ROW_H
    vals = {C_PERIOD: r.period, C_ROUTE: r.leg.route, C_FLT: r.leg.flight_no, C_AC: r.leg.aircraft, C_CFG: f'={_sheet_ref(sheet)}${L(C_CFG)}${first}', C_NOTE: ' / '.join(r.notes)}
    vals.update(_agg_formulas(row, first, last, sheet, blank='0'))
    if not r.cask.available:
        for c in (C_VAR, C_TOTCOST, C_CM, C_OP, C_CMR, C_OPR, C_PC, C_PCM, C_POP):
            vals[c] = None
    for i in range(2, LAST_COL):
        v = vals.get(i)
        c = ws.cell(row, i, v if v not in ('', None) else None)
        c.border = cell_border(i)
        c.font = F_BODY
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True) if i == C_NOTE else CENTER
        if i == C_LF:
            c.number_format = '0.0%;-0.0%;'
        elif i in (C_CMR, C_OPR):
            c.number_format = PCT
        elif i == C_AR:
            c.number_format = '#,##0;-#,##0;'
        elif i == C_RT:
            c.number_format = MONEY
        elif i in (C_CM, C_OP):
            c.number_format = PL_K
        elif i in (C_PCM, C_POP):
            c.number_format = PL_E
        elif i in (C_PR, C_PC):
            c.number_format = MONEY_E
        elif C_PAX <= i < C_NOTE:
            c.number_format = MONEY_K
    fill = None if r.cask.level == 'EXACT' else WARN_FILL if r.cask.available else NA_FILL
    if fill:
        for i in range(2, LAST_COL):
            ws.cell(row, i).fill = fill

def _write_scenario(ws, row: int, sc: Scenario, links: dict | None=None, scenario_desc: dict | None=None) -> tuple[int, int]:
    desc = (scenario_desc or {}).get(sc.name)
    _text(ws, row, 2, f'■ {sc.name}' + (f' - {desc}' if desc else ''), F_SECT)
    row += 1
    row = _write_header(ws, row)
    first = row
    for r in sc.results:
        link = links.get((sc.name, r.item)) if links else None
        if link:
            _write_link_row(ws, row, r, *link)
        else:
            _write_row(ws, row, r)
        row += 1
    last = row - 1
    if sc.no_total:
        missing = sum((1 for r in sc.results if not r.revenue_ok))
        if missing:
            _text(ws, row, 2, f'     ※ L/F·A/R 미확보 {missing}건 - 파란 칸에 값을 넣으면 자동 계산됩니다', F_NOTE)
            row += 1
        return (row + 1, 0)
    _write_total(ws, row, first, last)
    missing = sum((1 for r in sc.results if not r.revenue_ok))
    if missing:
        c = ws.cell(row, C_NOTE, f'※ L/F·A/R 미확보 {missing}건 - 파란 칸에 값을 넣으면 자동 계산됩니다')
        c.font = F_NOTE
        c.alignment = LEFT
    return (row + 2, row)

def _sheet_title(ws, title: str, main: bool=False):
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[2].height = 36.75
    t = ws.cell(2, 2, title)
    t.font = F_TITLE_W if main else F_TITLE
    if main:
        t.fill = TITLE_FILL
    t.alignment = CENTER
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=C_NOTE)

def _write_banner(ws, row: int, text: str) -> int:
    ws.row_dimensions[row].height = ROW_H
    for c in range(2, C_NOTE + 1):
        ws.cell(row, c).fill = BANNER_FILL
    t = ws.cell(row, 2, text)
    t.font = F_BANNER_W
    t.alignment = LEFT
    return row + 1

def _write_assumptions(ws, row: int, assumptions: list[str]) -> int:
    _text(ws, row, 2, '■ 산출기준', F_SECT)
    row += 1
    for a in assumptions:
        _text(ws, row, 2, a if a.startswith(' ') else f'     - {a}')
        row += 1
    return row
TOP_ASSUMPTION_PREFIXES = ('기간 :', '비용 :', '수입 :', '운항횟수 :', '조건 :')

def _trim_top_assumptions(assumptions: list[str]) -> list[str]:
    return [a for a in assumptions if a.startswith(TOP_ASSUMPTION_PREFIXES)]

def _set_widths(ws):
    ws.column_dimensions['A'].width = 1.33
    for i, (_, _, w) in enumerate(COLUMNS, start=2):
        ws.column_dimensions[L(i)].width = w
SUMMARY_WIDTHS = {'A': 1.375, 'B': 24.125, 'C': 11.25, 'D': 8.875, 'E': 7.875, 'F': 4.5, 'G': 8.875, 'H': 6.5, 'I': 8.0, 'J': 7.0, 'K': 6.0, 'L': 4.25, 'M': 7.0, 'P': 7.625, 'R': 8.375, 'S': 8.5, 'T': 5.625, 'V': 7.0, 'X': 41.875}

def _set_summary_widths(ws):
    for col, w in SUMMARY_WIDTHS.items():
        ws.column_dimensions[col].width = w
_BAD_SHEET = set('\\/*?:[]')

def _sheet_name(route: str, used: set) -> str:
    name = ''.join((ch for ch in route if ch not in _BAD_SHEET))[:31] or '노선'
    base, k = (name, 2)
    while name in used:
        name = f'{base[:27]}({k})'
        k += 1
    used.add(name)
    return name
BASIS_LAYOUT = [('기간', '기간', C_PERIOD, C_ROUTE), ('구분', '구분', C_FLT, C_CFG), ('L/F·A/R·화물수입 근거', '근거', C_RT, C_TOTREV), ('비고', '비고', C_VAR, C_OPR)]

def _write_basis_table(ws, row: int, rows: list[dict]) -> int:
    _text(ws, row, 2, '     - 월별 적용기준')
    row += 1
    ws.row_dimensions[row].height = HEAD_H
    for head, _, a, b in BASIS_LAYOUT:
        for c in range(a, b + 1):
            x = ws.cell(row, c)
            x.fill, x.font, x.alignment, x.border = (HEAD_FILL, F_HEAD_W, CENTER, BOX)
        ws.cell(row, a, head)
        if b > a:
            ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
    row += 1
    first = row
    same = {key: len({d.get(key) for d in rows}) == 1 for _, key, _, _ in BASIS_LAYOUT}
    for i, d in enumerate(rows):
        ws.row_dimensions[row].height = ROW_H
        for _, key, a, b in BASIS_LAYOUT:
            merge_down = key in ('기간', '비고') and same[key] and (len(rows) > 1)
            for c in range(a, b + 1):
                x = ws.cell(row, c)
                x.font, x.border, x.alignment = (F_BODY, BOX, CENTER)
            if merge_down and i > 0:
                continue
            ws.cell(row, a, d.get(key) or None)
            if merge_down:
                ws.merge_cells(start_row=row, start_column=a, end_row=row + len(rows) - 1, end_column=b)
            elif b > a:
                ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
        row += 1
    return row

def _write_route_sheet(wb, route: str, scenarios: list[Scenario], assumptions: list[str], title: str, used: set, item_labels: dict | None=None, basis_rows: list | None=None, scenario_desc: dict | None=None) -> dict:
    name = _sheet_name(route, used)
    ws = wb.create_sheet(name)
    _sheet_title(ws, title)
    row = _write_assumptions(ws, 4, assumptions)
    if basis_rows:
        row = _write_basis_table(ws, row, basis_rows)
    row += 1
    _text(ws, row, 2, '■ 월별 수지', F_SECT)
    row += 1

    def block(rows_, total: bool) -> tuple[int, int]:
        nonlocal row
        row = _write_header(ws, row)
        a = row
        for r in rows_:
            _write_row(ws, row, r)
            row += 1
        b = row - 1
        if total:
            _write_total(ws, row, a, b)
            missing = sum((1 for r in rows_ if not r.revenue_ok))
            if missing:
                c = ws.cell(row, C_NOTE, f'※ L/F·A/R 미확보 {missing}개월 - 파란 칸에 값을 넣으면 자동 계산됩니다')
                c.font = F_NOTE
                c.alignment = LEFT
            row += 1
        return (a, b)
    links = {}
    labels = item_labels or {}
    for sc in scenarios:
        mine = [r for r in sc.results if r.leg.route == route]
        if not mine:
            continue
        items = list(dict.fromkeys((r.item for r in mine)))
        local = list(dict.fromkeys((d for d in (labels.get(i) for i in items) if d)))
        desc = ' + '.join(local) if local else (scenario_desc or {}).get(sc.name)
        _text(ws, row, 2, f'○ {sc.name}' + (f' - {desc}' if desc else ''), F_SECT)
        row += 1
        if len(items) > 1 and (not sc.no_total):
            for k, item in enumerate(items):
                rs = [x for x in mine if x.item == item]
                head = labels.get(item) or rs[0].leg.aircraft
                _text(ws, row, 2, f'   - {head}', F_SECT)
                row += 1
                a, b = block(rs, total=True)
                links[sc.name, item] = (name, a, b)
                row += 1
        else:
            ordered = [x for item in items for x in mine if x.item == item]
            a, b = block(ordered, total=not sc.no_total)
            pos = a
            for item in items:
                n = sum((1 for x in mine if x.item == item))
                links[sc.name, item] = (name, pos, pos + n - 1)
                pos += n
            row += 1
    _set_widths(ws)
    return links

def write_excel(path: Path | str, title: str, scenarios: list[Scenario], assumptions: list[str], narrative: str, monthly: list[Scenario] | None=None, route_assumptions: dict | None=None, period_label: str='', item_labels: dict | None=None, route_tables: dict | None=None, scenario_desc: dict | None=None, match_rows: list | None=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = '수지비교'
    _sheet_title(ws, title, main=True)
    links = None
    if monthly:
        links, used = ({}, {'수지비교'})
        routes = list(dict.fromkeys((r.leg.route for sc in monthly for r in sc.results)))
        for route in routes:
            ra = (route_assumptions or {}).get(route) or assumptions
            links.update(_write_route_sheet(wb, route, monthly, ra, f'{route} 월별 수지 ({period_label})' if period_label else route, used, item_labels, (route_tables or {}).get(route), scenario_desc))
    live_all = [sc for sc in scenarios if sc.results]
    totalable = [sc for sc in live_all if not sc.no_total]
    scratch = wb.create_sheet('_layout_scratch')
    summary_h = _summary_table(scratch, 10, totalable, [1] * len(totalable), scenario_desc) - 10
    spans = []
    for i, sc in enumerate(live_all):
        r0 = 10 + i * 500
        r_end, tr = _write_scenario(scratch, r0, sc, links, scenario_desc)
        spans.append((r_end - r0, None if not tr else tr - r0))
    wb.remove(scratch)
    row = _write_banner(ws, 4, '1. 요약')
    row += 1
    top_assumptions = _trim_top_assumptions(assumptions) if monthly else assumptions
    row = _write_assumptions(ws, row, top_assumptions)
    row += 1
    summary_start = row
    scenario_start = summary_start + summary_h + 1 + 1
    cursor = scenario_start
    total_rows_all = []
    for consumed, offset in spans:
        total_rows_all.append(cursor + offset if offset is not None else 0)
        cursor += consumed
    totalable_total_rows = [tr for sc, tr in zip(live_all, total_rows_all) if not sc.no_total]
    row = _summary_table(ws, summary_start, totalable, totalable_total_rows, scenario_desc)
    row = _write_banner(ws, row, '2. 수지비교')
    row += 1
    for sc in live_all:
        row, _ = _write_scenario(ws, row, sc, links, scenario_desc)
    row += 1
    row = _sensitivity_table(ws, row, totalable)
    body = [l for l in _strip_summary_numbers(_strip_basis(narrative)).splitlines() if l.strip() != DISCLAIMER]
    while body and (not body[-1].strip()):
        body.pop()
    for line in body:
        _text(ws, row, 2, line)
        row += 1
    row += 1
    _text(ws, row, 2, DISCLAIMER, F_NOTE)
    row += 1
    _set_summary_widths(ws)
    if match_rows:
        _write_match_sheet(wb, match_rows)
    wb.active = 0
    path = Path(path)
    wb.save(path)
    return path
MATCH_COLS = [('노선', 14, None), ('(안)', 8, None), ('대상일', 11, 'yy.mm.dd'), ('요일', 5, None), ('대상 구분', 20, None), ('가져온 날', 30, None), ('사유', 52, None), ('편수', 6, '0.0'), ('공급석', 8, '#,##0'), ('수송석', 8, '#,##0'), ('L/F', 7, '0.0%'), ('A/R', 9, '#,##0'), ('예외', 5, None)]
MATCH_HOL_FILL = PatternFill('solid', fgColor='F8D7D3')

def _write_match_sheet(wb, rows: list) -> None:
    ws = wb.create_sheet('실적대응')
    _text(ws, 1, 1, '일자 매칭 실적대응 - 운항일마다 가져온 과거 실적 (예외 = 같은 요일 기본 규칙이 아닌 날)', F_SECT)
    for i, (head, width, _) in enumerate(MATCH_COLS, 1):
        c = ws.cell(3, i, head)
        c.fill, c.font, c.alignment, c.border = (HEAD_FILL, F_HEAD_W, CENTER, BOX)
        ws.column_dimensions[L(i)].width = width
    for r, d in enumerate(rows, 4):
        for i, (head, _, fmt) in enumerate(MATCH_COLS, 1):
            c = ws.cell(r, i, d.get(head))
            c.font, c.border = (F_BODY, BOX)
            if fmt:
                c.number_format = fmt
        if d.get('예외'):
            hol = '연휴' in d['대상 구분'] or ('대응' in d['사유'] and '→' not in d['사유'])
            for i in (5, 6, 7):
                ws.cell(r, i).fill = MATCH_HOL_FILL if hol else WARN_FILL
    ws.freeze_panes = 'D4'
    ws.auto_filter.ref = f'A3:{L(len(MATCH_COLS))}{len(rows) + 3}'
SUM_TABLE_COLS = [('총수입', C_PR), ('총비용', C_PC), ('한계이익', C_PCM), ('영업이익', C_POP)]
SENS_LAYOUT = [('노선', 2, 2), ('기종', 3, 3), ('유가 10c당', 4, 4), ('환율 10원당', 5, 6), ('총비용 대비', 7, 7)]

def _sens_cell(ws, row, a, b, value, fmt=None, head=False):
    c = ws.cell(row, a, value)
    c.font = F_HEAD_W if head else F_BODY
    if head:
        c.fill = HEAD_FILL
    c.alignment, c.border = (CENTER, BOX)
    if fmt:
        c.number_format = fmt
    for x in range(a + 1, b + 1):
        d = ws.cell(row, x)
        d.border = BOX
        if head:
            d.fill = HEAD_FILL
    if b > a:
        ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
    return c

def _sensitivity_table(ws, row: int, live) -> int:
    rows, seen = ([], set())
    for sc in live:
        for r in sc.results:
            key = (r.leg.route, r.leg.aircraft)
            if r.d_fuel and key not in seen:
                seen.add(key)
                rows.append(r)
    if not rows:
        return row
    _text(ws, row, 2, '■ 참고사항', F_SECT)
    row += 1
    _text(ws, row, 2, ' 유가·환율 민감도 (1왕복 변동비 증감, 천원)', F_SECT)
    row += 1
    ws.row_dimensions[row].height = HEAD_H
    for head, a, b in SENS_LAYOUT:
        _sens_cell(ws, row, a, b, head, head=True)
    row += 1
    for r in rows:
        ws.row_dimensions[row].height = ROW_H
        share = r.d_fuel * 10 / r.total_cost if r.total_cost else None
        vals = [r.leg.route, r.leg.aircraft, r.d_fuel * 10, r.d_fx * 10, share]
        fmts = [None, None, MONEY_K, MONEY_K, '0.0%']
        for (head, a, b), v, fmt in zip(SENS_LAYOUT, vals, fmts):
            _sens_cell(ws, row, a, b, v, fmt)
        row += 1
    _text(ws, row, 2, '     ※ 변동비는 유가·환율에 대해 1차식이라 위 증감액은 구간 무관 고정값입니다.', F_NOTE)
    return row + 2

def _summary_table(ws, row: int, live, total_rows, scenario_desc: dict | None=None) -> int:
    if not live or not total_rows or len(live) != len(total_rows):
        return row
    two = len(live) == 2
    heads = ['구분'] + [sc.name for sc in live] + (['차이'] if two else [])
    _text(ws, row, 2, '■ 총수지 요약 (억원)', F_SECT)
    row += 1
    if scenario_desc:
        for sc in live:
            desc = scenario_desc.get(sc.name)
            if desc:
                _text(ws, row, 2, f'     {sc.name} = {desc}')
                row += 1
    ws.row_dimensions[row].height = ROW_H
    for i, h in enumerate(heads, start=2):
        c = ws.cell(row, i, h)
        c.font, c.fill, c.alignment, c.border = (F_HEAD_W, HEAD_FILL, CENTER, BOX)
    row += 1
    for label, col in SUM_TABLE_COLS:
        ws.row_dimensions[row].height = ROW_H
        vals = [label] + [f'={L(col)}{tr}' for tr in total_rows]
        if two:
            vals.append(f'={L(col)}{total_rows[1]}-{L(col)}{total_rows[0]}')
        for i, v in enumerate(vals, start=2):
            c = ws.cell(row, i, v)
            c.font = F_SUM if i == 2 else F_BODY
            c.alignment, c.border = (CENTER, BOX)
            if i > 2:
                c.number_format = PL_E if col in (C_PCM, C_POP) or (two and i == 5) else MONEY_E
        row += 1
    for label, col in (('한계이익률', C_CMR), ('영업이익률', C_OPR)):
        ws.row_dimensions[row].height = ROW_H
        vals = [label] + [f'={L(col)}{tr}' for tr in total_rows]
        if two:
            vals.append(f'={L(col)}{total_rows[1]}-{L(col)}{total_rows[0]}')
        for i, v in enumerate(vals, start=2):
            c = ws.cell(row, i, v)
            c.font = F_SUM if i == 2 else F_BODY
            c.alignment, c.border = (CENTER, BOX)
            if i > 2:
                c.number_format = PCT
        row += 1
    ws.row_dimensions[row].height = ROW_H
    for i, v in enumerate(['운항횟수(왕복)'] + [f'={L(C_RT)}{tr}' for tr in total_rows] + ([f'={L(C_RT)}{total_rows[1]}-{L(C_RT)}{total_rows[0]}'] if two else []), start=2):
        c = ws.cell(row, i, v)
        c.font = F_SUM if i == 2 else F_BODY
        c.alignment, c.border = (CENTER, BOX)
        if i > 2:
            c.number_format = MONEY
    return row + 2

def _eok(v):
    return f'{v / 100000000.0:,.0f}억원'

def _man(v):
    return f'{v / 10000.0:,.0f}만원'

def _contribution_lines(base: Scenario, alt: Scenario) -> list[str]:

    def per_route(sc):
        d = {}
        for r in sc.results:
            if not r.revenue_ok:
                continue
            e = d.setdefault(r.leg.route, {'cm': 0.0, 'op': 0.0, 'legs': []})
            e['cm'] += r.period_total('contribution') or 0
            e['op'] += r.period_total('operating_profit') or 0
            e['legs'].append(r)
        return d
    b, a = (per_route(base), per_route(alt))
    routes = list(dict.fromkeys(list(b) + list(a)))
    if not routes:
        return []
    diffs = []
    for rt in routes:
        bb, aa = (b.get(rt, {'cm': 0, 'op': 0}), a.get(rt, {'cm': 0, 'op': 0}))
        diffs.append((rt, aa['cm'] - bb['cm'], aa['op'] - bb['op']))
    diffs.sort(key=lambda x: -abs(x[1]))
    out = ['', '○ 차이 요인 (노선별 한계이익 기여, 총수지 기준)']
    for rt, dcm, dop in diffs:
        sign = '+' if dcm >= 0 else ''
        legs_b = ', '.join((f'{r.leg.aircraft} {r.leg.round_trips:,.0f}왕복' for r in b.get(rt, {}).get('legs', []))) or '-'
        legs_a = ', '.join((f'{r.leg.aircraft} {r.leg.round_trips:,.0f}왕복' for r in a.get(rt, {}).get('legs', []))) or '-'
        out.append(f'     - {rt} : {sign}{_eok(dcm)}   ({legs_b} → {legs_a})')
    top, tdcm, _ = diffs[0]
    total = sum((d[1] for d in diffs))
    if total and abs(tdcm / total) >= 0.5:
        out.append(f'     → 차이의 대부분({_eok(tdcm)} / 전체 {_eok(total)})이 {top} 에서 발생')
    cmp_lines = []
    for rt in routes:
        seen = {}
        for r in b.get(rt, {}).get('legs', []) + a.get(rt, {}).get('legs', []):
            seen.setdefault(r.leg.aircraft, r)
        if len(seen) < 2:
            continue
        ranked = sorted(seen.items(), key=lambda kv: -(kv[1].contribution or 0))
        txt = ' / '.join((f'{ac} {_man(r.contribution)}({r.seats}석)' for ac, r in ranked))
        gap = (ranked[0][1].contribution or 0) - (ranked[-1][1].contribution or 0)
        cmp_lines.append(f'     - {rt} : {txt}  → {ranked[0][0]} 우위 {_man(gap)}/왕복')
    if cmp_lines:
        out.append('')
        out.append('○ 동일 노선 기종별 1왕복 한계이익')
        out += cmp_lines
    return out

def _by_aircraft_lines(sc) -> list:
    rows = [r for r in sc.results if r.total_cost is not None]
    if not rows:
        return ['', '○ 기종별 1왕복 비용 : 산출 가능한 기종이 없습니다']
    route = rows[0].leg.route
    out = ['', f'○ {route} 기종별 1왕복 비용 (천원)', '']
    out.append(f"     {'기종':<9}{'좌석':>6}{'변동비':>12}{'총비용':>12}{'좌석당총비용':>14}{'한계이익':>12}{'영업이익':>12}")
    for r in sorted(rows, key=lambda x: x.total_cost):
        seats = r.seats_offered / 2
        cm = '-' if r.contribution is None else f'{r.contribution / 1000:>12,.0f}'
        op = '-' if r.operating_profit is None else f'{r.operating_profit / 1000:>12,.0f}'
        mark = r.leg.aircraft + ('*' if r.cask.is_proxy else '')
        out.append(f'     {mark:<9}{seats:>6,.0f}{r.variable_cost / 1000:>12,.0f}{r.total_cost / 1000:>12,.0f}{r.total_cost / seats / 1000:>14,.1f}{cm:>12}{op:>12}')
    n_proxy = sum((1 for r in rows if r.cask.is_proxy))
    if n_proxy:
        out.append('')
        out.append(f'     * {n_proxy}/{len(rows)}개 기종은 이 노선 CASK 미보유로 B/T 유사노선 단가를 쓴 「BT기준 산정」 값입니다.')
        out.append('       기종 간 우열을 이 값으로 확정하지 말고 경영기획 원가 확인 필요.')
    cheap = min(rows, key=lambda x: x.total_cost)
    unit = min(rows, key=lambda x: x.total_cost / x.seats_offered)
    out.append('')
    out.append(f'     → 1왕복 총비용이 가장 낮은 기종 : {cheap.leg.aircraft} ({_won(cheap.total_cost)})')
    if unit.leg.aircraft != cheap.leg.aircraft:
        out.append(f'     → 좌석당 총비용이 가장 낮은 기종 : {unit.leg.aircraft} (좌석당 {unit.total_cost / (unit.seats_offered / 2) / 1000:,.1f}천원) - 수요가 받쳐주면 대형기가 단위원가에서 유리')
    prof = [r for r in rows if r.operating_profit is not None]
    if prof:
        best = max(prof, key=lambda x: x.operating_profit)
        out.append(f'     → 영업이익이 가장 큰 기종 : {best.leg.aircraft} ({_won(best.operating_profit)})  ※해당 노선 과거실적 L/F·A/R 적용 기준')
    out.append('     ※ 기종끼리 합산하는 값이 아니므로 합계행은 만들지 않았습니다.')
    return out

def _won(v):
    return f'{v / 10000.0:,.0f}만원'
DISCLAIMER = '※ 본 문구는 자동 생성 초안이며, 최종 표현 및 결론은 검수 후 확정 필요.'

def _drop_block(narrative: str, head: str) -> str:
    out, skip = ([], False)
    for line in narrative.splitlines():
        t = line.strip()
        if t.startswith(head):
            skip = True
            continue
        if skip:
            if t.startswith('○') or t.startswith('■'):
                skip = False
            else:
                continue
        out.append(line)
    return '\n'.join(out)

def _strip_summary_numbers(narrative: str) -> str:
    drop_heads = ('○ 예상 수지', '○ 기존(안) 대비')
    narrative = _drop_block(narrative, '○ 유가·환율 민감도')
    out, skip = ([], False)
    for line in narrative.splitlines():
        t = line.strip()
        if any((t.startswith(h) for h in drop_heads)):
            skip = True
            continue
        if skip:
            if t.startswith('-') or t.startswith('- ') or t.startswith('· '):
                continue
            if t.startswith('→') or t.startswith('※'):
                out.append(line)
                skip = False
                continue
            if t.startswith('○') or t.startswith('■'):
                skip = False
            elif not t:
                continue
        if not skip:
            out.append(line)
    while out and (not out[0].strip()):
        out.pop(0)
    return '\n'.join(out)

def _strip_basis(narrative: str) -> str:
    out, skip = ([], False)
    for line in narrative.splitlines():
        t = line.strip()
        if t.startswith('○ 산출기준'):
            skip = True
            continue
        if skip and (t.startswith('○') or t.startswith('■')):
            skip = False
        if not skip:
            out.append(line)
    while out and (not out[0].strip()):
        out.pop(0)
    return '\n'.join(out)

def _caveat_lines(proxies, missing, no_rev) -> list:
    out = []
    out.append('     - 기재 변경에 따른 고정비 분산 효과는 미반영 (해당 노선 배부 몫만 산출)')
    if no_rev:
        out.append(f"     - 과거실적 미보유로 L/F·A/R 미확보 → 수지 산출 제외 : {', '.join(sorted(set(no_rev)))}")
    if proxies:
        out.append('     - 아래 노선은 해당 기종의 CASK 미보유로 대체 산정한 값이며(동일노선 자매기종 또는 B/T 유사노선 단가), 실제 원가와 차이 발생 가능 (경영기획 원가 확정 필요)')
        out.append(f"        · {', '.join(sorted(set(proxies)))}")
    if missing:
        out.append('     - 아래 노선은 CASK 및 대체 기준이 모두 없어 산출 불가 (경영기획 원가 요청 필요)')
        out.append(f"        · {', '.join(sorted(set(missing)))}")
    return out

def build_narrative(scenarios: list[Scenario], assumptions: list[str], scenario_desc: dict | None=None) -> str:
    lines = ['○ 산출기준']
    lines += [f'     - {a}' if not a.startswith(' ') else a for a in assumptions]
    proxies, missing, no_rev = ([], [], [])
    for sc in scenarios:
        for r in sc.results:
            tag = f'{r.leg.route}({r.leg.aircraft})'
            if not r.cask.available:
                missing.append(tag)
            elif r.cask.is_proxy:
                proxies.append(tag)
            if not r.revenue_ok:
                no_rev.append(tag)
    live = [sc for sc in scenarios if sc.results]
    if len(live) == 1 and live[0].no_total:
        lines += _by_aircraft_lines(live[0])
    elif len(live) == 1:
        sc = live[0]
        lines.append('')
        lines.append(f'○ 예상 수지 (총수지, 억원)')
        if not any((r.revenue_ok for r in sc.results)):
            cost = sum(((sens_tot(r) or 0) * r.leg.round_trips for r in sc.results))
            var = sum(((sens_var(r) or 0) * r.leg.round_trips for r in sc.results))
            lines.append('     - 총수입 : L/F·A/R 미확보 → 산출 불가')
            lines.append(f'     - 총비용 : {_eok(cost)} (변동비 {_eok(var)}) ※판매수수료·전산비 등 수입연동비 제외')
            lines.append('     - 한계이익·영업이익 : 수입이 없어 판단 불가')
            lines.append("     → 노선 시트 파란 칸에 L/F·A/R 을 넣거나, 입력에 'L/F 85% A/R 250000' 을 붙여 다시 돌리면 손익까지 계산됨.")
            lines.append('')
            lines.append('○ 유의사항')
            lines += _caveat_lines(proxies, missing, no_rev)
            lines.append('')
            lines.append(DISCLAIMER)
            return '\n'.join(lines)
        for label, attr in (('총수입', 'total_revenue'), ('총비용', 'total_cost'), ('한계이익', 'contribution'), ('영업이익', 'operating_profit')):
            v, _ = sc.total_partial(attr)
            lines.append(f'     - {label} : {_eok(v)}')
        rev, _ = sc.total_partial('total_revenue')
        cm, _ = sc.total_partial('contribution')
        op, _ = sc.total_partial('operating_profit')
        if rev:
            lines.append(f'     - 한계이익률 {cm / rev:.1%} / 영업이익률 {op / rev:.1%}')
        if cm > 0 and op < 0:
            lines.append('     → 한계이익은 흑자이나 고정비 배부 후 영업이익은 적자. 기재를 이미 보유·운영 중이라면 운항 자체는 기여하나,')
            lines.append('        노선 단독으로 고정비까지 회수하지는 못하는 구조임.')
        elif cm <= 0:
            lines.append('     → 한계이익이 적자 - 운항할수록 손실이 커지는 구조로, 운임·탑승률 개선 없이는 운항 재검토 필요.')
        else:
            lines.append('     → 고정비까지 회수하는 흑자 구조.')
    blind = [sc.name for sc in live if not any((r.revenue_ok for r in sc.results))]
    if len(live) == 2 and blind:
        lines.append('')
        lines.append('○ 수지 비교')
        if scenario_desc:
            for sc in live:
                d = scenario_desc.get(sc.name)
                if d:
                    lines.append(f'     {sc.name} = {d}')
        lines.append(f"     - {', '.join(blind)} 은 L/F·A/R 미확보로 수입·손익을 산출하지 못해 총수지 비교를 하지 않았습니다.")
        for sc in live:
            cost = sum(((sens_tot(r) or 0) * r.leg.round_trips for r in sc.results))
            lines.append(f'     - {sc.name} 총비용 : {_eok(cost)}' + (' ※수입연동비 제외' if sc.name in blind else ''))
        lines.append("     → 노선 시트 파란 칸에 L/F·A/R 을 넣거나 입력에 'L/F 85% A/R 250000' 을 붙여 다시 돌리면 비교됨.")
    elif len(live) == 2:
        base, alt = scenarios
        lines.append('')
        lines.append(f'○ {base.name} 대비 {alt.name} 수지 변동 (총수지, 억원)')
        if scenario_desc:
            for sc in (base, alt):
                d = scenario_desc.get(sc.name)
                if d:
                    lines.append(f'     {sc.name} = {d}')
        diff = {}
        for label, attr in (('총수입', 'total_revenue'), ('총비용', 'total_cost'), ('한계이익', 'contribution'), ('영업이익', 'operating_profit')):
            b, _ = base.total_partial(attr)
            a, _ = alt.total_partial(attr)
            diff[label] = a - b
            lines.append(f"     - {label} : {_eok(b)} → {_eok(a)} ({('+' if a - b >= 0 else '')}{_eok(a - b)})")
        op, cm = (diff['영업이익'], diff['한계이익'])
        if op > 0:
            lines.append(f'     → 변경(안)의 영업이익이 기존(안) 대비 약 {_eok(abs(op))} 개선되어 변경(안) 추진이 유리한 것으로 분석됨.')
        elif op < 0:
            lines.append(f'     → 기존(안)의 영업이익이 변경(안) 대비 약 {_eok(abs(op))} 우위로, 기존 운영안 유지가 유리한 것으로 분석됨.')
        else:
            lines.append('     → 양 (안) 간 영업이익 차이가 유의미하지 않음.')
        if cm * op < 0:
            better_cm, better_op = ('변경(안)', '기존(안)') if cm > 0 else ('기존(안)', '변경(안)')
            fc_b, _ = base.total_partial('fixed_cost')
            fc_a, _ = alt.total_partial('fixed_cost')
            lines.append('')
            lines.append('     ※ 두 지표의 결론이 엇갈림 - 판단 주의')
            lines.append(f'        · 한계이익 기준 : {better_cm} 우위 ({_eok(abs(cm))})')
            lines.append(f'        · 영업이익 기준 : {better_op} 우위 ({_eok(abs(op))})')
            lines.append(f'        · 영업이익 차이는 편수 변동에 따른 고정비 배부액 변화({_eok(abs(fc_a - fc_b))})가 주된 요인임.')
            lines.append('          리스료·감가상각 등은 감편해도 실제로 소멸하지 않고 타 노선으로 재배부되므로,')
            lines.append('          기재를 이미 보유·운영 중이라면 한계이익 기준 판단이 타당함.')
            lines.append('          단, 감편으로 확보한 기재를 타 노선에 돌린다면 투입 노선의 한계이익이')
            lines.append('          감편한 편의 한계이익보다 커야 이득임(기회비용).')
            lines.append('          이때 기재는 시간 자원이므로 1왕복당이 아니라 B/T 시간당 한계이익으로 비교할 것.')
            lines.append('          - 1왕복 한계이익이 비슷해도 B/T 13.4h 노선 1편을 빼면 B/T 5.2h 노선 2편을')
            lines.append('            돌릴 수 있어 시간당으로는 2배 이상 벌어진다.')
            lines.append('          - 슬롯·정비·승무원 제약으로 그 가동시간을 실제로 옮길 수 있을 때만 성립.')
    if len(live) == 2 and (not blind):
        lines += _contribution_lines(scenarios[0], scenarios[1])
    rows = [r for sc in live for r in sc.results if r.d_fuel]
    if rows:
        lines.append('')
        lines.append('○ 유가·환율 민감도 (1왕복 변동비 증감, 천원)')
        lines.append(f"     {'노선':<16}{'기종':<7}{'유가 10c당':>12}{'환율 10원당':>13}{'총비용 대비':>13}")
        seen = set()
        for r in rows:
            key = (r.leg.route, r.leg.aircraft)
            if key in seen:
                continue
            seen.add(key)
            share = f'{r.d_fuel * 10 / r.total_cost:>12.1%}' if r.total_cost else f"{'-':>12}"
            lines.append(f'     {r.leg.route:<16}{r.leg.aircraft:<7}{r.d_fuel * 10 / 1000:>12,.0f}{r.d_fx * 10 / 1000:>13,.0f}{share}')
        lines.append('     ※ 변동비는 유가·환율에 대해 1차식이라 위 증감액은 구간 무관 고정값입니다.')
    lines.append('')
    lines.append('○ 유의사항')
    lines += _caveat_lines(proxies, missing, no_rev)
    lines.append('')
    lines.append(DISCLAIMER)
    return '\n'.join(lines)

def sens_var(r):
    return r.variable_cost if r.variable_cost is not None else r.direct_var

def sens_tot(r):
    if r.total_cost is not None:
        return r.total_cost
    return None if r.direct_var is None else r.direct_var + r.fix_const

def sens_partial(cells) -> bool:
    return any((r is not None and r.total_cost is None for r in cells.values()))

def sens_missing_note(cells) -> str:
    rs = [r for r in cells.values() if r is not None and r.total_cost is None]
    if not rs:
        return ''
    items = ['판매·가맹점 수수료', '전산비']
    rev_fix = any((r.fix_coef_rev for r in rs))
    if rev_fix:
        items.append('간접고정비 중 운송수입 배부분')
    note = 'L/F·A/R 미확보 → ' + ' · '.join(items) + ' 제외한 비용'
    if rev_fix:
        note += ' (간접고정비를 편수·B/T 로 배부하면 그것까지 포함된다)'
    return note

def write_sensitivity(path, title: str, fx_list: list, fuel_list: list, blocks, assumptions: list, base=None):
    if isinstance(blocks, dict):
        blocks = [('', blocks)]
    blocks = [(b[0], b[1], b[2] if len(b) > 2 else '') for b in blocks]
    combos = [(fx, fu) for fx in fx_list for fu in fuel_list]
    multi = len(blocks) > 1
    routes = list(dict.fromkeys((rt for _, _, rt in blocks)))
    split = multi and len(routes) > 1 and all(routes)
    wide = max(len(fx_list), len(combos) if multi else 0)
    wb = Workbook()
    ws = wb.active
    row = 0

    def start_sheet(sheet, name, heading, cols, info):
        nonlocal ws, row
        ws = sheet
        ws.title = name
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 1.33
        ws.row_dimensions[2].height = 36.75
        t = ws.cell(2, 2, heading)
        t.font = F_TITLE
        t.alignment = CENTER
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=max(3, 2 + cols))
        row = 4
        if info:
            _text(ws, row, 2, '■ 산출기준', F_SECT)
            row += 1
            for a in info:
                _text(ws, row, 2, f'     - {a}')
                row += 1
            row += 1

    def head_cell(r, c, v, fmt=None):
        x = ws.cell(r, c, v)
        x.font, x.fill, x.alignment, x.border = (F_HEAD_W, HEAD_FILL, CENTER, BOX)
        if fmt:
            x.number_format = fmt
        return x

    def compare(label, get, fmt):
        nonlocal row
        _text(ws, row, 2, f'○ {label}', F_SECT)
        row += 1
        ws.row_dimensions[row].height = ROW_H
        head_cell(row, 2, '노선 · 기종')
        for j, (fx, fu) in enumerate(combos):
            head_cell(row, 3 + j, f'{fx:,.0f}원 / {fu:,.0f}c')
        row += 1
        for name, cells, _ in blocks:
            ws.row_dimensions[row].height = ROW_H
            head_cell(row, 2, name)
            for j, k in enumerate(combos):
                r_ = cells.get(k)
                c2 = ws.cell(row, 3 + j, None if r_ is None else get(r_))
                c2.font, c2.alignment, c2.border = (F_BODY, CENTER, BOX)
                c2.number_format = fmt
            row += 1
        row += 1

    def grid(cells, label, get, fmt, note=''):
        nonlocal row
        _text(ws, row, 2, f'○ {label}', F_SECT)
        row += 1
        if note:
            _text(ws, row, 2, f'     {note}', F_NOTE)
            row += 1
        ws.row_dimensions[row].height = ROW_H
        head_cell(row, 2, '유가(c) / 환율(원)')
        for j, fx in enumerate(fx_list):
            head_cell(row, 3 + j, fx, FX_UNIT)
        row += 1
        for fu in fuel_list:
            ws.row_dimensions[row].height = ROW_H
            head_cell(row, 2, fu, FUEL_UNIT)
            for j, fx in enumerate(fx_list):
                r_ = cells.get((fx, fu))
                v = None if r_ is None else get(r_)
                c2 = ws.cell(row, 3 + j, v)
                c2.font, c2.alignment, c2.border = (F_BODY, CENTER, BOX)
                c2.number_format = fmt
            row += 1
        row += 1

    def metrics(cells):
        yield ('1왕복 변동비 (천원)', lambda r: sens_var(r), MONEY_K, '')
        yield ('1왕복 총비용 (천원)', lambda r: sens_tot(r), MONEY_K, '')
        b = cells.get(base) if base is not None else None
        bt0 = None if b is None else sens_tot(b)
        if bt0 is not None:
            yield (f'기준({base[0]:,.0f}원 / {base[1]:,.0f}USC) 대비 총비용 증감 (천원)', lambda r: sens_tot(r) - bt0 if sens_tot(r) is not None else None, PL_K, '＋면 기준보다 비싸다')
        if any((r is not None and r.operating_profit is not None for r in cells.values())):
            yield ('1왕복 영업이익 (천원)', lambda r: r.operating_profit, PL_K, '')
            yield ('1왕복 한계이익 (천원)', lambda r: r.contribution, PL_K, '')

    def widths(cols, wide_b):
        for i in range(2, 3 + cols):
            ws.column_dimensions[L(i)].width = 16 if multi else 14
        if wide_b:
            ws.column_dimensions['B'].width = 20

    def grids(items, heading_each):
        nonlocal row
        for name, cells, _ in items:
            if heading_each:
                _text(ws, row, 2, f'■ {name}', F_SECT)
                row += 1
            for label, get, fmt, note in metrics(cells):
                grid(cells, label, get, fmt, note)
    start_sheet(ws, '요약' if split else '민감도', title, wide, assumptions)
    if multi:
        compare('노선별 1왕복 총비용 (천원)', sens_tot, MONEY_K)
        compare('노선별 1왕복 변동비 (천원)', sens_var, MONEY_K)
        if any((r is not None and r.operating_profit is not None for _, cs, _ in blocks for r in cs.values())):
            compare('노선별 1왕복 영업이익 (천원)', lambda r: r.operating_profit, PL_K)
    if not split:
        grids(blocks, multi)
        widths(wide, multi)
    else:
        widths(wide, True)
        _text(ws, row, 2, '     노선별 환율 x 유가 그리드는 노선 시트에 있습니다 : ' + ' · '.join(routes), F_NOTE)
        used = {'요약'}
        for rt in routes:
            items = [b for b in blocks if b[2] == rt]
            name = _sheet_name(rt, used)
            start_sheet(wb.create_sheet(), name, f'{rt} 환율 x 유가 민감도', len(fx_list), [])
            grids(items, True)
            widths(len(fx_list), True)
    path = Path(path)
    wb.save(path)
    return path
