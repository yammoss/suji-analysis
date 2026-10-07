# -*- coding: utf-8 -*-
from __future__ import annotations
from datetime import date, datetime
from pathlib import Path
import openpyxl
from openpyxl import Workbook
from .paths import DATA_DIR
NAME = '비용기준 DATA.xlsx'
SHEET = '월별비용'
HEAD = ['노선', '기종', '년월', '구분', '출처', '편수(왕복)', '공급석', '수송석', '운송수입', '부대수입', '화물수입', '변동비', '고정비', '총비용']
KEYS = ('rt', 'seats', 'pax', 'pax_rev', 'anc', 'cargo', 'var', 'fix', 'cost')
DAY_SHEET = '운항일'
DAY_HEAD = ['노선', '기종', '운항일', '편수(OW)', '공급석']

def path() -> Path | None:
    if DATA_DIR is None:
        return None
    p = DATA_DIR / NAME
    return p if p.exists() else None

def write(target: Path, rows: list, days: list | None=None) -> Path:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(SHEET)
    ws.append(HEAD)
    for route, ac, y, m, kind, label, v in rows:
        ws.append([route, ac or '', datetime(y, m, 1), kind, label, *(None if v.get(k) is None else round(float(v[k]), 1) for k in KEYS)])
    if days:
        wd = wb.create_sheet(DAY_SHEET)
        wd.append(DAY_HEAD)
        for route, ac, d, fc, seats in days:
            wd.append([route, ac, datetime(d.year, d.month, d.day), fc, seats])
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target

class CostBook:

    def __init__(self, p: Path | None=None):
        self.path = p or path()
        self.available = False
        self.months: list[tuple[int, int]] = []
        self._by: dict[tuple[str, int, int], dict] = {}
        self._plan: dict[tuple[str, int, int], dict] = {}
        self._label: dict[tuple[int, int], tuple[str, str]] = {}
        self._days: dict[str, dict[date, dict[str, list]]] = {}
        if self.path is None or not self.path.exists():
            return
        try:
            wb = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
            ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
            day_rows = list(wb[DAY_SHEET].iter_rows(values_only=True)) if DAY_SHEET in wb.sheetnames else []
            wb.close()
        except Exception:
            return
        for r in day_rows[1:]:
            if not r or len(r) < 5 or (not r[0]) or (not hasattr(r[2], 'year')):
                continue
            d = r[2].date() if hasattr(r[2], 'date') else r[2]
            self._days.setdefault(str(r[0]).strip(), {}).setdefault(d, {})[str(r[1] or '')] = [float(r[3] or 0), float(r[4] or 0)]
        for r in rows[1:]:
            if not r or not r[0] or (not hasattr(r[2], 'year')):
                continue
            route, y, m, kind, label = (str(r[0]).strip(), r[2].year, r[2].month, r[3], r[4])
            store = self._plan if str(kind or '').strip() == '계획' else self._by
            cur = store.setdefault((route, y, m), {k: 0.0 for k in KEYS})
            for i, k in enumerate(KEYS, start=5):
                if i < len(r) and isinstance(r[i], (int, float)):
                    cur[k] += float(r[i])
            if store is self._by:
                self._label[y, m] = (str(kind or ''), str(label or ''))
        self.months = sorted(self._label)
        self.available = bool(self._by)

    def get(self, route: str, y: int, m: int) -> dict | None:
        return self._by.get((route, y, m))

    def fill_capacity(self, act) -> int:
        if not self._days or getattr(act, '_cap_day', None):
            return 0
        for route, days in self._days.items():
            cap = act._cap_day.setdefault(route, {})
            cap_ac = act._cap_day_ac.setdefault(route, {})
            for d, per_ac in days.items():
                cap[d] = [sum((v[0] for v in per_ac.values())), sum((v[1] for v in per_ac.values()))]
                cap_ac[d] = {a: list(v) for a, v in per_ac.items()}
                for a, (fc, seats) in per_ac.items():
                    e = act._cap.setdefault((route, a, d.year, d.month), [0.0, 0.0])
                    e[0] += fc
                    e[1] += seats
        return len(self._days)

    def plan(self, route: str, y: int, m: int) -> dict | None:
        return self._plan.get((route, y, m))

    def kind(self, y: int, m: int) -> str:
        return self._label.get((y, m), ('', ''))[0]

    def label_of(self, y: int, m: int) -> str:
        return self._label.get((y, m), ('', ''))[1]

    def covers(self, y: int, m: int) -> bool:
        return (y, m) in self._label

    @property
    def note(self) -> str:
        if not self.available:
            return ''
        ms = self.months
        return f'비용기준 DATA ({ms[0][0] % 100}.{ms[0][1]:02d}~{ms[-1][0] % 100}.{ms[-1][1]:02d})'
