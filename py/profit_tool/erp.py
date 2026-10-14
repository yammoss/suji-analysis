# -*- coding: utf-8 -*-
from __future__ import annotations
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .paths import DATA_DIR, PPT_DIR
NAME = 'ERP 실적 DATA.xlsx'
MIL = 1000000.0
KR_AIRPORTS = {'ICN', 'GMP', 'PUS', 'CJU', 'TAE', 'CJJ', 'KWJ', 'MWX', 'RSU', 'USN', 'YNY', 'KUV', 'WJU', 'HIN', 'KPO', 'YEC'}
ERP_KEYS = ['flights', 'lf_w', 'rev', 'cargo', 'misc', 'total_rev', 'var', 'fixed', 'cost', 'margin', 'op']
BI_KEYS = ['flights', 'supply', 'pax', 'rev', 'OB_flights', 'OB_supply', 'OB_pax', 'OB_rev', 'IB_flights', 'IB_supply', 'IB_pax', 'IB_rev']

def route_key(name):
    t = re.sub('\\(.*?\\)', '', str(name or '')).replace('V.V', '').strip().upper()
    parts = [p.strip() for p in re.split('[-~/\\s]+', t) if p.strip()]
    if len(parts) == 1 and re.fullmatch('[A-Z]{6}', parts[0]):
        parts = [parts[0][:3], parts[0][3:]]
    return frozenset(parts[:2]) if len(parts) >= 2 else None
KR_ORDER = ['ICN', 'GMP', 'PUS', 'TAE', 'CJJ', 'KWJ', 'MWX', 'RSU', 'USN', 'YNY', 'KUV', 'WJU', 'HIN', 'KPO', 'YEC', 'CJU']

def route_name(k) -> str:
    rank = lambda x: (0, KR_ORDER.index(x)) if x in KR_ORDER else (1, x)
    ab = sorted(k, key=rank)
    return f'{ab[0]}-{ab[-1]}'

def path() -> Path | None:
    p = DATA_DIR / NAME if DATA_DIR else None
    return p if p and p.exists() else None

def _sources() -> list[Path]:
    if PPT_DIR is None:
        return []
    out = [PPT_DIR / 'RAW DATA' / 'RAW.xlsx']
    try:
        from .weekly import find_files
        out += [Path(f) for f in find_files()]
    except Exception:
        pass
    leg = PPT_DIR.parent / '6수요' / 'raw'
    out += [Path(p) for p in leg.glob('*년 실적*.xlsx') if not p.name.startswith('~$')] if leg.is_dir() else []
    return [p for p in out if p.exists()]

def stale(target: Path) -> bool:
    src = _sources()
    if not src:
        return False
    return not target.exists() or target.stat().st_mtime < max((p.stat().st_mtime for p in src))

def _week_label(text: str) -> str:
    m = re.search('(\\d{1,2})\\s*월\\s*(\\d)\\s*주차', str(text or '').split('_')[-1])
    return f'{int(m.group(1))}월 {m.group(2)}주차 추정' if m else '추정'

def _add_weekly(erp: dict) -> dict:
    try:
        from .weekly import Weekly
        wk = Weekly()
    except Exception:
        return {}
    if not wk.available:
        return {}
    last = max((ym for ym, _ in erp), default=0)
    labels = {}
    for (route, y, m), v in wk._by.items():
        ym = y * 100 + m
        k = route_key(route)
        if ym <= last or not k or (not v.get('rt')):
            continue
        fc = v['rt'] * 2
        total = v.get('rev') or (v.get('pax_rev') or 0) + (v.get('anc') or 0) + (v.get('cargo') or 0)
        e = erp.setdefault((ym, k), {'sched': set(), 'lf_w': 0.0})
        e.setdefault('sched', set()).add('부정기노선' if 'CHTR' in str(route).upper() else '정기노선')
        for kk, val in (('flights', fc), ('rev', v.get('pax_rev')), ('cargo', v.get('cargo')), ('misc', v.get('anc')), ('total_rev', total), ('var', v.get('var')), ('fixed', v.get('fix')), ('cost', v.get('cost')), ('margin', total - (v.get('var') or 0)), ('op', total - (v.get('cost') or 0))):
            e[kk] = e.get(kk, 0.0) + (val or 0.0)
        if v.get('seats'):
            e['lf_w'] += (v.get('pax') or 0) / v['seats'] * fc
        labels[ym] = _week_label(wk.label_of(y, m))
    return labels

def build(target_dir: Path) -> str:
    if PPT_DIR is None or not (PPT_DIR / 'src' / 'route_history.py').exists():
        return 'ERP 실적 DATA : PPT 자동화 폴더를 찾지 못해 만들지 않았습니다'
    sys.path.insert(0, str(PPT_DIR / 'src'))
    try:
        import route_history as RH
        erp, est = RH.load_erp()
        leg, leg_name = RH.load_leg()
    except Exception as e:
        return f'ERP 실적 DATA : 원본을 읽지 못했습니다 ({e}) - 기존 파일을 그대로 둡니다'
    finally:
        sys.path.remove(str(PPT_DIR / 'src'))
    weekly = _add_weekly(erp)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('ERP')
    ws.append(['노선', '년월', '기준', '정기/부정기'] + ERP_KEYS)
    for (ym, k), v in sorted(erp.items(), key=lambda x: (route_name(x[0][1]), x[0][0])):
        ws.append([route_name(k), datetime(ym // 100, ym % 100, 1), weekly.get(ym) or ('추정' if ym in est else 'ERP'), '·'.join(sorted((x for x in v.get('sched', set()) if x)))] + [round(float(v.get(kk) or 0), 6 if kk == 'lf_w' else 1) for kk in ERP_KEYS])
    wb2 = wb.create_sheet('BI')
    wb2.append(['노선', '년월'] + BI_KEYS)
    for (ym, k), v in sorted(leg.items(), key=lambda x: (route_name(x[0][1]), x[0][0])):
        wb2.append([route_name(k), datetime(ym // 100, ym % 100, 1)] + [round(float(v.get(kk) or 0), 1) for kk in BI_KEYS])
    info = wb.create_sheet('정보')
    info.append(['BI 원본', leg_name or ''])
    info.append(['만든 시각', datetime.now().strftime('%Y-%m-%d %H:%M')])
    target_dir = Path(target_dir)
    tmp = Path(__file__).resolve().parent.parent / 'output' / NAME
    tmp.parent.mkdir(exist_ok=True)
    wb.save(tmp)
    dst = target_dir / NAME
    try:
        import shutil
        shutil.copy2(tmp, dst)
    except OSError as e:
        return f'ERP 실적 DATA : 공유폴더에 쓰지 못했습니다 ({e})'
    yms = sorted({ym for ym, _ in erp})
    wk_txt = ' · 주차 추정 ' + ', '.join((f'{ym // 100 % 100}.{ym % 100:02d}({lab})' for ym, lab in sorted(weekly.items()))) if weekly else ''
    return f'ERP 실적 DATA : ERP {len(erp):,}행 ({yms[0] // 100 % 100}.{yms[0] % 100}~{yms[-1] // 100 % 100}.{yms[-1] % 100}, 추정 {len(est)}개월{wk_txt}) · BI {len(leg):,}행 → {dst.name}'

def load(p: Path | None=None):
    p = p or path()
    if p is None:
        return None
    wb = load_workbook(p, read_only=True, data_only=True)
    erp, est, leg = ({}, {}, {})
    for r in wb['ERP'].iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or (not hasattr(r[1], 'year')):
            continue
        ym = r[1].year * 100 + r[1].month
        k = route_key(r[0])
        e = {kk: float(r[4 + i] or 0) for i, kk in enumerate(ERP_KEYS) if 4 + i < len(r)}
        e['sched'] = set((x for x in str(r[3] or '').split('·') if x))
        erp[ym, k] = e
        if r[2] and str(r[2]) != 'ERP':
            est[ym] = str(r[2])
    for r in wb['BI'].iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or (not hasattr(r[1], 'year')):
            continue
        ym = r[1].year * 100 + r[1].month
        leg[ym, route_key(r[0])] = {kk: float(r[2 + i] or 0) for i, kk in enumerate(BI_KEYS) if 2 + i < len(r)}
    leg_name = ''
    if '정보' in wb.sheetnames:
        for r in wb['정보'].iter_rows(values_only=True):
            if r and r[0] == 'BI 원본':
                leg_name = str(r[1] or '')
    wb.close()
    return (erp, est, leg, leg_name)
HDR = PatternFill('solid', fgColor='666058')
SUB = PatternFill('solid', fgColor='F2E9E3')
GRP = {'손익 (ERP, 백만원)': PatternFill('solid', fgColor='9C5B4B'), '운송 (BI)': PatternFill('solid', fgColor='666058')}
THIN = Side(style='thin', color='D9D9D9')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COLS = [('년월', None, None, None), ('기준', None, None, None), ('정기/부정기', None, None, None), ('운항편수', '손익 (ERP, 백만원)', lambda p, l: p.get('flights'), '#,##0'), ('탑승률', '손익 (ERP, 백만원)', lambda p, l: p['lf_w'] / p['flights'] if p.get('flights') else None, '0.0%'), ('운송수입', '손익 (ERP, 백만원)', lambda p, l: p.get('rev', 0) / MIL, '#,##0'), ('화물수입', '손익 (ERP, 백만원)', lambda p, l: p.get('cargo', 0) / MIL, '#,##0'), ('부대수입', '손익 (ERP, 백만원)', lambda p, l: p.get('misc', 0) / MIL, '#,##0'), ('총수입', '손익 (ERP, 백만원)', lambda p, l: p.get('total_rev', 0) / MIL, '#,##0'), ('변동비', '손익 (ERP, 백만원)', lambda p, l: p.get('var', 0) / MIL, '#,##0'), ('고정비', '손익 (ERP, 백만원)', lambda p, l: p.get('fixed', 0) / MIL, '#,##0'), ('총비용', '손익 (ERP, 백만원)', lambda p, l: p.get('cost', 0) / MIL, '#,##0'), ('한계이익', '손익 (ERP, 백만원)', lambda p, l: p.get('margin', 0) / MIL, '#,##0;[Red]-#,##0'), ('한계이익률', '손익 (ERP, 백만원)', lambda p, l: p['margin'] / p['total_rev'] if p.get('total_rev') else None, '0.0%;[Red]-0.0%'), ('영업이익', '손익 (ERP, 백만원)', lambda p, l: p.get('op', 0) / MIL, '#,##0;[Red]-#,##0'), ('영업이익률', '손익 (ERP, 백만원)', lambda p, l: p['op'] / p['total_rev'] if p.get('total_rev') else None, '0.0%;[Red]-0.0%'), ('운항편수', '운송 (BI)', lambda p, l: l.get('flights'), '#,##0'), ('공급석', '운송 (BI)', lambda p, l: l.get('supply'), '#,##0'), ('수송석', '운송 (BI)', lambda p, l: l.get('pax'), '#,##0'), ('L/F', '운송 (BI)', lambda p, l: l['pax'] / l['supply'] if l.get('supply') else None, '0.0%'), ('A/R(원)', '운송 (BI)', lambda p, l: l['rev'] / l['pax'] if l.get('pax') else None, '#,##0'), ('O/B 수송석', '운송 (BI)', lambda p, l: l.get('OB_pax'), '#,##0'), ('O/B L/F', '운송 (BI)', lambda p, l: l['OB_pax'] / l['OB_supply'] if l.get('OB_supply') else None, '0.0%'), ('O/B A/R', '운송 (BI)', lambda p, l: l['OB_rev'] / l['OB_pax'] if l.get('OB_pax') else None, '#,##0'), ('I/B 수송석', '운송 (BI)', lambda p, l: l.get('IB_pax'), '#,##0'), ('I/B L/F', '운송 (BI)', lambda p, l: l['IB_pax'] / l['IB_supply'] if l.get('IB_supply') else None, '0.0%'), ('I/B A/R', '운송 (BI)', lambda p, l: l['IB_rev'] / l['IB_pax'] if l.get('IB_pax') else None, '#,##0')]

def _sum(dicts):
    out = {}
    for d in dicts:
        for k, v in d.items():
            if isinstance(v, (int, float)):
                out[k] = out.get(k, 0.0) + v
    return out

def write_route(wb, code, erp, est, leg, leg_name):
    k = route_key(code)
    if not k:
        return None
    yms = sorted({ym for ym, kk in erp if kk == k} | {ym for ym, kk in leg if kk == k})
    if not yms:
        return None
    name = route_name(k)
    ws = wb.create_sheet(name[:31])
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, '%s 월별 실적 (%d.%d ~ %d.%d)' % (name, yms[0] // 100 % 100, yms[0] % 100, yms[-1] // 100 % 100, yms[-1] % 100)).font = Font(bold=True, size=13, color='666058')
    ws.cell(2, 1, "※ 손익 = ERP (24·25년 'YYYY ERP DATA', 26년 '실제 DATA', 추정 = 추정실적 파일 · 'm월 n주차 추정' = 주차별 추정실적 보고) · 운송 = BI 방향별 실적 (%s) · 금액 백만원 · 24·25년 영업이익은 그해 비용 배부 기준" % (leg_name or '-')).font = Font(size=9, color='8C8C8C')
    r0 = 4
    groups = {}
    for j, (h, g, _fn, _fmt) in enumerate(COLS, 1):
        if g:
            groups.setdefault(g, []).append(j)
        c = ws.cell(r0 + 1, j, h)
        c.fill, c.font, c.border = (HDR, Font(bold=True, color='FFFFFF', size=10), BOX)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for g, js in groups.items():
        ws.merge_cells(start_row=r0, start_column=js[0], end_row=r0, end_column=js[-1])
        c = ws.cell(r0, js[0], g)
        c.fill, c.font = (GRP[g], Font(bold=True, color='FFFFFF', size=10))
        c.alignment = Alignment(horizontal='center')
    for j in (1, 2, 3):
        ws.merge_cells(start_row=r0, start_column=j, end_row=r0 + 1, end_column=j)
        c = ws.cell(r0, j, COLS[j - 1][0])
        c.fill, c.font, c.border = (HDR, Font(bold=True, color='FFFFFF', size=10), BOX)
        c.alignment = Alignment(horizontal='center', vertical='center')
    row = r0 + 2

    def put(label, basis, sched, p, l, fill=None, bold=False):
        nonlocal row
        vals = [label, basis, sched] + [fn(p, l) if fn else None for _h, _g, fn, _f in COLS[3:]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v)
            c.border = BOX
            c.font = Font(size=10, bold=bold)
            if fill:
                c.fill = fill
            fmt = COLS[j - 1][3]
            if fmt and isinstance(v, (int, float)):
                c.number_format = fmt
        row += 1
    for y in sorted({ym // 100 for ym in yms}):
        ps, ls = ([], [])
        for ym in [x for x in yms if x // 100 == y]:
            p = erp.get((ym, k), {})
            l = leg.get((ym, k), {})
            ps.append(p)
            ls.append(l)
            sched = '·'.join(sorted((x.replace('노선', '') for x in p.get('sched', set()) if x))) if p else ''
            basis = (est.get(ym) if isinstance(est, dict) else '추정' if ym in est else None) or 'ERP' if p else '-'
            put('%d.%02d' % (y % 100, ym % 100), basis, sched, p, l)
        put('%d년 계' % y, '', '', _sum(ps), _sum(ls), fill=SUB, bold=True)
    widths = [9, 6, 10] + [10] * (len(COLS) - 3)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 2, 4)
    return (name, yms)
