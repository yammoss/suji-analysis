# -*- coding: utf-8 -*-
from __future__ import annotations
import argparse
import sys
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from profit_tool.actuals import Actuals
from profit_tool.dataset import Dataset
from profit_tool.engine import Engine, Leg
from profit_tool.period import Period
from profit_tool.report import BOX, CENTER, EOK, F_BODY, F_HEAD_W, F_SECT, F_TITLE_W, HEAD_FILL, LEFT, MONEY, MONEY_K, PCT, PL_E, PL_K, TITLE_FILL, TOTAL_FILL
BASE = Path(__file__).resolve().parent

def eprint(*a):
    print(*a, file=sys.stderr)

def build_rows(ds: Dataset, act: Actuals, period: Period, eng: Engine, years_back: int) -> list[dict]:
    rows = []
    for route, plan in ds.plan_by_route.items():
        info = ds.routes.get(route)
        if not info:
            continue
        for ac, fc_arr in plan.get('by_type', {}).items():
            fc_sum = sum((fc_arr[i] for i in eng.mi if i < len(fc_arr)))
            rt = fc_sum / 2.0
            if rt < 0.5:
                continue
            r = act.for_period(route, period, years_back=years_back, use_plan=True)
            name = ds.aircraft.get(ac, {}).get('name', ac)
            c = act.cargo_revenue(route, name, period, years_back=years_back)
            leg = Leg(flight_no='', route=route, aircraft=ac, round_trips=rt, lf=r.lf if r else None, ar=r.ar if r else None, cargo_rt=c.per_round_trip)
            res = eng.compute(leg)
            rows.append(dict(route=route, region=info.get('region', ''), aircraft=ac, seats=res.seats, distance=info.get('distance_km'), round_trips=rt, lf=res.leg.lf, ar=res.leg.ar, cask_level=res.cask.level, cask_note=res.cask.note, revenue_ok=res.revenue_ok, pax_revenue=res.pax_revenue, total_revenue=res.total_revenue, total_cost=res.total_cost, contribution=res.contribution, operating_profit=res.operating_profit, cm_rate=res.contribution_margin, op_rate=res.operating_margin, period_revenue=res.period_total('total_revenue'), period_cost=res.period_total('total_cost'), period_cm=res.period_total('contribution'), period_op=res.period_total('operating_profit'), lf_src=r.short if r else 'L/F·A/R 미확보'))
    return rows
COLS = [('순위', 5), ('노선', 16), ('대노선', 12), ('기종', 7), ('좌석', 6), ('기간왕복', 9), ('L/F', 7), ('A/R', 9), ('총수입(1왕복,천원)', 13), ('총비용(1왕복,천원)', 13), ('한계이익(1왕복,천원)', 13), ('영업이익(1왕복,천원)', 13), ('한계이익률', 9), ('영업이익률', 9), ('기간 영업이익(억원)', 13), ('비고', 42)]

def write_report(rows: list[dict], period_label: str, fixed_alloc: str, out: Path):
    live = [r for r in rows if r['revenue_ok']]
    missing = [r for r in rows if not r['revenue_ok']]
    live.sort(key=lambda r: r['period_op'])
    wb = Workbook()
    ws = wb.active
    ws.title = '전노선 헬스체크'
    ws.sheet_view.showGridLines = False
    ws.merge_cells('B2:P2')
    t = ws.cell(2, 2, f"전 노선 헬스체크 ({period_label}, 배부기준: {('편수·B/T' if fixed_alloc == 'volume' else '운송수입')})")
    t.font, t.fill, t.alignment = (F_TITLE_W, TITLE_FILL, CENTER)
    ws.row_dimensions[2].height = 32
    tot_rev = sum((r['period_revenue'] or 0 for r in live))
    tot_cost = sum((r['period_cost'] or 0 for r in live))
    tot_op = sum((r['period_op'] or 0 for r in live))
    loss_n = sum((1 for r in live if (r['period_op'] or 0) < 0))
    row = 4
    ws.cell(row, 2, '■ 요약').font = F_SECT
    row += 1
    for label in (f'     - 대상 : {period_label} 사업계획 편성 노선 x 기종 {len(live)}건 (L/F·A/R 미확보 {len(missing)}건 제외)', f'     - 네트워크 합계 : 총수입 {tot_rev / EOK:,.0f}억 / 총비용 {tot_cost / EOK:,.0f}억 / 영업이익 {tot_op / EOK:+,.0f}억', f'     - 영업이익 적자 노선 : {loss_n}개 / 흑자 {len(live) - loss_n}개', '     - 이 표는 시나리오 비교가 아니라 현재 편성된 사업량 그대로의 스냅샷입니다 (운항횟수 = W26 비용파일 사업량, 스케줄 재입력 아님)'):
        ws.cell(row, 2, label).font = F_BODY
        row += 1
    row += 1
    ws.cell(row, 2, '■ 노선 x 기종별 손익 (영업이익 낮은 순)').font = F_SECT
    row += 1
    hdr_row = row
    for i, (h, w) in enumerate(COLS, start=2):
        c = ws.cell(hdr_row, i, h)
        c.font, c.fill, c.alignment, c.border = (F_HEAD_W, HEAD_FILL, CENTER, BOX)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hdr_row].height = 24
    row += 1
    for rank, r in enumerate(live, start=1):
        note = r['lf_src'] + (f" / {r['cask_note']}" if r['cask_note'] else '')
        vals = [rank, r['route'], r['region'], r['aircraft'], r['seats'], r['round_trips'], r['lf'], r['ar'], r['total_revenue'], r['total_cost'], r['contribution'], r['operating_profit'], r['cm_rate'], r['op_rate'], r['period_op'] / EOK, note]
        fmts = [None, None, None, None, MONEY, '#,##0', '0.0%', MONEY, MONEY_K, MONEY_K, PL_K, PL_K, PCT, PCT, PL_E, None]
        for i, (v, fmt) in enumerate(zip(vals, fmts), start=2):
            c = ws.cell(row, i, v)
            c.font, c.border = (F_BODY, BOX)
            c.alignment = LEFT if i in (3, 17) else CENTER
            if fmt:
                c.number_format = fmt
        ws.row_dimensions[row].height = 18
        row += 1
    tot = row
    ws.cell(tot, 2, '네트워크 합계').font = Font(name=F_BODY.name, bold=True)
    for c in range(2, len(COLS) + 2):
        ws.cell(tot, c).fill, ws.cell(tot, c).border = (TOTAL_FILL, BOX)
    ws.cell(tot, 7, f'=SUM(G{hdr_row + 1}:G{tot - 1})').number_format = '#,##0'
    op_cell = ws.cell(tot, 16, tot_op / EOK)
    op_cell.number_format, op_cell.font = (PL_E, Font(name=F_BODY.name, bold=True))
    ws.cell(tot, 17, f'총수입 {tot_rev / EOK:,.0f}억 / 총비용 {tot_cost / EOK:,.0f}억').font = F_BODY
    if missing:
        row = tot + 2
        ws.cell(row, 2, '■ L/F·A/R 미확보 (제외됨)').font = F_SECT
        row += 1
        ws.cell(row, 2, '     · ' + ', '.join((f"{r['route']}({r['aircraft']})" for r in missing))).font = F_BODY
    ws.freeze_panes = ws.cell(hdr_row + 1, 3).coordinate
    ws.auto_filter.ref = f'B{hdr_row}:{get_column_letter(len(COLS) + 1)}{tot - 1}'
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out

def main():
    ap = argparse.ArgumentParser(description='전 노선 헬스체크 - 현재 편성 그대로 노선별 손익 스캔')
    ap.add_argument('--period', default='W26')
    ap.add_argument('--fixed-alloc', choices=('revenue', 'volume'), default='volume')
    ap.add_argument('--actuals-back', type=int, default=1)
    ap.add_argument('--fx', type=float, default=None)
    ap.add_argument('--fuel', type=float, default=None)
    ap.add_argument('--out', default=None)
    args = ap.parse_args()
    ds = Dataset()
    act = Actuals()
    period = Period.parse(args.period)
    eng = Engine(ds, period, fx=args.fx, fuel=args.fuel, fixed_alloc=args.fixed_alloc)
    eprint(f'■ 대상 기간 : {period.label}  (배부기준: {args.fixed_alloc})')
    rows = build_rows(ds, act, period, eng, args.actuals_back)
    eprint(f"   {len(rows)}건 계산 완료 (L/F·A/R 미확보 {sum((1 for r in rows if not r['revenue_ok']))}건 포함)")
    out = Path(args.out) if args.out else BASE / 'output' / f'전노선_헬스체크_{datetime.now():%y%m%d}.xlsx'
    write_report(rows, period.label, args.fixed_alloc, out)
    eprint(f'\n[저장] {out}')
if __name__ == '__main__':
    main()
