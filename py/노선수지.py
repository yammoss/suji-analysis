# -*- coding: utf-8 -*-
from __future__ import annotations
import argparse
import sys
from datetime import datetime, timedelta
from profit_tool.actuals import Actuals
from profit_tool.dataset import Dataset
from profit_tool.engine import Engine, Leg, Scenario
from profit_tool.period import Period
from profit_tool.plan27 import Plan
from openpyxl.utils import get_column_letter as L
from profit_tool.report import BOX, CENTER, F_BODY, F_HEAD_W, F_SUM, HEAD_FILL, LEFT, MONEY_K, PCT, PL_K, ROW_H, _sheet_name, _text, build_narrative, write_excel
from profit_tool.schedule import month_spans, parse_schedule
from profit_tool.costbook import CostBook
from profit_tool.mgmt import Mgmt
from profit_tool.weekly import Weekly
BASE = __import__('pathlib').Path(__file__).resolve().parent
EOK = 100000000

def eprint(*a):
    print(*a, file=sys.stderr)

def month_rows(ds, act, route, period, sched, ac_hint, alloc, item=0, wk=None, mg=None, cb=None, pl=None):
    for lo, hi in month_spans(period):
        y, m = (lo.year, lo.month)
        mp = Period.parse(f'{y % 100}.{m:02d}')
        wk_scaled = {}
        eng, pool_note = _engine_for(ds, mp, alloc)
        agg = act._agg.get((route, y, m))
        mix = {a: v for (r, a, yy, mm), v in act._cargo.items() if r == route and yy == y and (mm == m) and v[1]}
        ext_row = ext_kind = ext_label = None
        if mg is not None and mg.available:
            ext_row = mg.get(route, y, m)
            if ext_row and ext_row.get('fc'):
                ext_row = dict(ext_row, rt=ext_row['fc'] / 2)
                ext_kind, ext_label = ('확정', f'확정실적({mg.short})')
            else:
                ext_row = None
        if ext_row is None and wk is not None and wk.available:
            got = wk.get(route, y, m)
            if got and got.get('rt'):
                ext_row, ext_kind = (got, '추정')
                ext_label = f'추정실적({wk.label_of(y, m)})'
        if ext_row is None and cb is not None and cb.available:
            got = cb.get(route, y, m)
            if got and got.get('rt'):
                ext_row, ext_kind = (got, cb.kind(y, m) or '추정')
                ext_label = f'{ext_kind}실적({cb.label_of(y, m)})'
        part = _daily_span(act, route, lo, hi)
        full_month = lo.day == 1 and (hi + timedelta(days=1)).month != m
        last_flown = act.day_range[1] if act.day_range else None
        running = bool(full_month and part and last_flown and ((y, m) == (last_flown.year, last_flown.month)) and (part[5] < hi))
        if running and ext_row and ext_row.get('rt'):
            part, mix = (None, {})
        if mix and (part or (agg and agg[0] > 0 and full_month)):
            clipped = part and (not full_month or part[5] < hi)
            month_fc = {a: v[1] for a, v in mix.items()}
            if clipped:
                fc_tot, seats, pax, rev, d0, d1 = part
                span_fc = act.capacity_span_ac(route, lo, hi) or month_fc
                share = fc_tot / max(sum(month_fc.values()), 1e-09)
                last = act.day_range[1] if act.day_range else d1
                running = full_month and (y, m) == (last.year, last.month)
                src = f'{d0:%y.%m.%d}~{d1:%m.%d} 실적' + (' (월 미완 - 그날까지)' if running else ' (운항일 기준)')
            else:
                seats, pax, rev = agg
                fc_tot, share, span_fc = (sum(month_fc.values()), 1.0, month_fc)
                src = f'{y % 100}.{m:02d}월 실적'
            rt = fc_tot / 2
            cargo_tot = sum((v[0] for v in mix.values())) * share
            legs = []
            tot_fc = max(sum(span_fc.values()), 1e-09)
            for name, fc in span_fc.items():
                code = next((c for c, d in ds.aircraft.items() if d['name'] == name), None)
                if code is None:
                    continue
                rt_i = rt * fc / tot_fc
                cargo_ac = mix.get(name, [0.0, 0.0])[0] * (fc / max(month_fc.get(name) or fc, 1e-09))
                legs.append((code, rt_i, cargo_ac / max(rt_i, 1e-09)))
            if not legs:
                legs = [(_main_ac(act, ds, route), rt, 0.0)]
        elif ext_row and ext_row.get('rt'):
            span_fc, _ = act.capacity_span(route, lo, hi)
            month_fc, _ = act.capacity_span(route, lo.replace(day=1), (lo.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1))
            if span_fc and month_fc:
                k = span_fc / month_fc
            else:
                days = (hi - lo).days + 1
                whole = ((lo.replace(day=28) + timedelta(days=4)).replace(day=1) - lo.replace(day=1)).days
                k = days / max(whole, 1)
            rt = ext_row['rt'] * k
            seats = (ext_row['seats'] or 0) * k
            pax = (ext_row['pax'] or 0) * k
            rev = (ext_row['pax_rev'] or 0) * k
            cargo_tot = (ext_row['cargo'] or 0) * k
            mixw = act.capacity(route, y, m)
            code = (next((c for c, d in ds.aircraft.items() if d['name'] == max(mixw, key=lambda x: mixw[x][0])), None) if mixw else None) or ac_hint or _main_ac(act, ds, route)
            legs = [(code, rt, cargo_tot / max(rt, 1e-09))]
            src = ext_label + (f' · 구간 {k:.0%}' if k < 0.999 else '')
        else:
            cap_fc, cap_seats = act.capacity_span(route, lo, hi)
            ac_seg = None
            plan = cb.plan(route, y, m) if cb is not None and cb.available and (not cap_fc) else None
            sched_note = ''
            if cap_fc:
                rt = cap_fc / 2
                sched_note = '확정 스케줄'
            elif plan and plan.get('rt'):
                days = (hi - lo).days + 1
                whole = (lo.replace(day=28) + timedelta(days=4)).replace(day=1) - lo.replace(day=1)
                rt = plan['rt'] * (days / max(whole.days, 1))
                sched_note = '확정 스케줄'
            elif pl is not None and pl.available and pl.volume(route, y, m):
                fleet = pl.fleet(route, y, m)
                days = (hi - lo).days + 1
                whole = ((lo.replace(day=28) + timedelta(days=4)).replace(day=1) - lo.replace(day=1)).days
                rt = sum(fleet.values()) / 2 * (days / max(whole, 1))
                ac_seg = ds.resolve_aircraft(max(fleet, key=fleet.get))
                sched_note = f'{pl.short} 사업량'
            elif sched:
                span = Period.parse(f'{lo:%y.%m.%d}~{hi:%y.%m.%d}')
                r = parse_schedule(sched, span)
                if r is None:
                    yield dict(y=y, m=m, empty=True, src=f"스케줄 '{sched}' 인식 실패")
                    continue
                rt = r.round_trips
                sched_note = sched
            else:
                seg, ac_seg = _plan_volume(ds, route, y, m)
                if not seg:
                    yield dict(y=y, m=m, empty=True, src='실적·확정 스케줄·사업량 없음 (--sched 로 스케줄을 주면 추정)')
                    continue
                days = (hi - lo).days + 1
                whole = ((lo.replace(day=28) + timedelta(days=4)).replace(day=1) - lo.replace(day=1)).days
                rt = seg / 2 * (days / max(whole, 1))
                sched_note = 'W26 파일 사업량'
            plan_lf = plan_ar = None
            if pl is not None and pl.available:
                plan_lf, plan_ar = pl.rate(route, y, m)
            a = act.for_period(route, mp, years_back=1, use_plan=True)
            if a is None and (not (plan_lf and plan_ar)):
                yield dict(y=y, m=m, empty=True, src='L/F·A/R 실적·계획 없음')
                continue
            mixc = act.capacity(route, y, m)
            code = ac_hint or (next((c for c, d in ds.aircraft.items() if d['name'] == max(mixc, key=lambda k2: mixc[k2][0])), None) if mixc else None) or (ac_seg if sched_note == 'W26 파일 사업량' and ac_seg else None) or _main_ac(act, ds, route)
            cargo = act.cargo_revenue(route, ds.aircraft[code]['name'], mp, years_back=1)
            span_ac = act.capacity_span_ac(route, lo, hi) or {} if cap_fc else {}
            mix_legs = []
            if not ac_hint and len(span_ac) > 1:
                tot_fc = max(sum(span_ac.values()), 1e-09)
                for name, fc in span_ac.items():
                    c2 = next((c for c, d in ds.aircraft.items() if d['name'] == name), None)
                    if c2 is not None:
                        mix_legs.append((c2, rt * fc / tot_fc))
            if len(mix_legs) > 1:
                seats = sum((ds.seats(c2) * r2 * 2 for c2, r2 in mix_legs))
                sched_note += ' (기종 ' + '+'.join((f"{ds.aircraft[c2]['name']} {r2:.0f}왕복" for c2, r2 in mix_legs)) + ')'
            else:
                mix_legs = [(code, rt)]
                seats = ds.seats(code) * rt * 2
            if plan_lf and plan_ar:
                lf_src = f'{pl.short} 목표'
                pax, rev = (seats * plan_lf, seats * plan_lf * plan_ar)
            else:
                lf_src = a.source_label
                pax, rev = (seats * a.lf, seats * a.lf * a.ar)
            legs = [(c2, r2, cargo.per_round_trip) for c2, r2 in mix_legs]
            cargo_tot = cargo.per_round_trip * rt
            src = f'{pl.short} 목표·사업량' if plan_lf and plan_ar and sched_note.startswith(pl.short) else lf_src + f' · {sched_note}'
        fx, fuel, idx_src = (eng.fx, eng.fuel, f'{y % 100}.{m:02d} INDEX')
        if ext_kind == '추정' and wk is not None and wk.available:
            srcw = wk.source(y, m)
            if srcw.get('fx'):
                fx, fuel, idx_src = (srcw['fx'], srcw.get('fuel') or fuel, srcw.get('label', idx_src))
        if ext_row and ext_row.get('rt'):
            wk_scaled = {a2: (ext_row.get(a2) or 0) / ext_row['rt'] for a2 in ('var', 'fix', 'anc')}
            src += f' · 비용 {ext_label}'
        lf = pax / seats if seats else 0.0
        ar = rev / pax if pax else 0.0
        cost = anc = var = 0.0
        results = []
        for code, rt_i, cargo_rt in legs:
            res = eng.compute(Leg(flight_no='', route=route, aircraft=code, round_trips=rt_i, lf=lf, ar=ar, cargo_rt=cargo_rt, rt_source=src))
            if wk_scaled:
                _apply_weekly(res, wk_scaled)
            res.item = item
            res.period = mp.label
            results.append(res)
            cost += res.total_cost * rt_i
            var += (res.variable_cost or 0) * rt_i
            anc += (res.ancillary_revenue or 0) * rt_i
        names = ' + '.join((f"{ds.aircraft[c]['name']}" for c, _, _ in legs))
        yield dict(y=y, m=m, empty=False, ac=names, ow=round(sum((r_ * 2 for _, r_, _ in legs))), results=results, legs=legs, seats=seats, pax=pax, lf=lf, ar=ar, rev=rev, anc=anc, cargo=cargo_tot, total_rev=rev + anc + cargo_tot, cost=cost, var=var, rt=sum((r_ for _, r_, _ in legs)), fallback=getattr(eng, 'season_fallback', False), pool_fix=bool(pool_note), weekly=bool(wk_scaled), cost_src=ext_kind or 'CASK', fx=fx, fuel=fuel, idx_src=idx_src, src=src if wk_scaled else src + ' · 비용 CASK(W26 파일' + (f', {pool_note.split()[-2]}' if pool_note else '') + f', {y % 100}.{m:02d} INDEX)')

def _apply_weekly(res, per_rt):
    for tgt_key, cur_attr, put_attr, const_attr in (('var', 'variable_cost', 'indirect_var', 'var_const'), ('fix', 'fixed_cost', 'indirect_fix', 'fix_const')):
        tgt = per_rt.get(tgt_key)
        cur = getattr(res, cur_attr)
        if tgt is None or cur is None:
            continue
        d = tgt - cur
        setattr(res, put_attr, (getattr(res, put_attr) or 0) + d)
        setattr(res, const_attr, (getattr(res, const_attr) or 0) + d)
    if res.ancillary_revenue is not None and per_rt.get('anc') is not None:
        d = per_rt['anc'] - res.ancillary_revenue
        res.ancillary_revenue += d
        res.anc_const += d
        if res.total_revenue is not None:
            res.total_revenue += d

def _engine_for(ds, mp, alloc):
    eng = Engine(ds, mp, fixed_alloc=alloc)
    if not getattr(eng, 'partial_only', False):
        return (eng, '')
    keys = set(mp.keys)
    here = next((i for i, lab in enumerate(ds.months) if f"20{lab.replace('.', '-')}" in keys), None)
    cands = [i for i in range(len(ds.months)) if i not in ds.partial_months]
    if here is None or not cands:
        return (eng, '')
    pick = min(cands, key=lambda i: abs(i - here))
    alt = Engine(ds, Period.parse(ds.months[pick]), fixed_alloc=alloc, fx=eng.fx, fuel=eng.fuel)
    return (alt, f'고정비 배부는 {ds.months[pick]} 단가')

def _plan_volume(ds, route, y, m):
    key = f'{y % 100:02d}.{m:02d}'
    if key not in ds.months:
        return (0.0, None)
    i = ds.months.index(key)
    plan = ds.plan_by_route.get(route)
    if not plan:
        return (0.0, None)
    seg = plan['total'][i] if i < len(plan['total']) else 0.0
    detail = (ds.plan_detail or {}).get(route) or {}
    best, best_fc = (None, 0.0)
    for code, v in detail.items():
        fc = (v.get('fc') or [0])[i] if i < len(v.get('fc') or []) else 0.0
        if fc > best_fc:
            best, best_fc = (code, fc)
    return (seg, best)

def _daily_span(act, route, lo, hi):
    days = act._day.get(route)
    if not days:
        return None
    got = [(d, v) for d, v in days.items() if lo <= d <= hi and v[1] > 0]
    if not got:
        return None
    tot = [sum((v[i] for _, v in got)) for i in range(4)]
    return tot + [min((d for d, _ in got)), max((d for d, _ in got))]

def _main_ac(act, ds, route):
    cnt = {}
    for (r, name, _, _), v in act._cargo.items():
        if r == route:
            cnt[name] = cnt.get(name, 0) + v[1]
    if not cnt:
        return 'B738'
    name = max(cnt, key=cnt.get)
    return next((c for c, d in ds.aircraft.items() if d['name'] == name), 'B738')

def _month_ranges(rows, kind):
    ms = sorted({(d['y'], d['m']) for d in rows if d.get('cost_src') == kind})
    if not ms:
        return ''
    out, start, prev = ([], ms[0], ms[0])
    for cur in ms[1:]:
        nxt = (prev[0] + (prev[1] == 12), prev[1] % 12 + 1)
        if cur != nxt:
            out.append((start, prev))
            start = cur
        prev = cur
    out.append((start, prev))
    f = lambda t: f'{t[0] % 100}.{t[1]:02d}'
    return ' · '.join((f(a) if a == b else f'{f(a)}~{f(b)}' for a, b in out))

def _cost_lines(live, alloc, mg, wk):
    alloc_txt = '편수·B/T' if alloc == 'volume' else '운송수입'
    out = []
    span = _month_ranges(live, '확정')
    if span:
        out.append(f'비용 : {span} 는 경영기획 노선별 확정실적' + (f' ({mg.short})' if mg is not None and mg.available else '') + ' 의 변동비·고정비 그대로')
    span = _month_ranges(live, '추정')
    if span:
        files = ' · '.join(dict.fromkeys((d['src'].split('(')[-1].rstrip(')').split(' · ')[0] for d in live if d.get('cost_src') == '추정')))
        out.append(f'비용 : {span} 는 주차별 추정실적({files}) 의 변동비·고정비 그대로 (그 주차 기준 환율·유가 INDEX 반영)')
    span = _month_ranges(live, 'CASK')
    if span:
        out.append(f'비용 : {span} 는 W26 비용추정용 파일의 기종별-노선별 CASK + 해당 월 환율·유가 INDEX, 간접고정비 {alloc_txt} 기준 배부')
    if not out:
        out.append(f'비용 : W26 비용추정용 파일 CASK + 해당 월 INDEX, 간접고정비 {alloc_txt} 배부')
    return out

def _fill(period, blocks, period_results, monthly_results, route_assumptions, alloc, ds):
    for item, (route, rows) in enumerate(blocks):
        got = [d for d in rows if not d['empty']]
        if not got:
            continue
        for d in got:
            monthly_results += d['results']
        rt = sum((d['rt'] for d in got))
        seats = sum((d['seats'] for d in got))
        pax = sum((d['pax'] for d in got))
        rev = sum((d['rev'] for d in got))
        cargo = sum((d['cargo'] for d in got))
        code = max({c: sum((r_ for cc, r_, _ in d['legs'] if cc == c)) for d in got for c, _, _ in d['legs']}.items(), key=lambda x: x[1])[0]
        eng, _ = _engine_for(ds, period, alloc)
        res = eng.compute(Leg(flight_no='', route=route, aircraft=code, round_trips=rt, lf=pax / seats if seats else None, ar=rev / pax if pax else None, cargo_rt=cargo / rt if rt else 0.0, rt_source=f"실적 {sum((d['ow'] for d in got)):,.0f}편(OW) = {rt:,.0f}왕복"))
        res.item = item
        res.period = period.label
        period_results.append(res)
        rt_r = sum((d['rt'] for d in got)) or 1e-09
        fx_r = sum(((d.get('fx') or 0) * d['rt'] for d in got)) / rt_r
        fu_r = sum(((d.get('fuel') or 0) * d['rt'] for d in got)) / rt_r
        route_assumptions.setdefault(route, []).extend([f'기간 : {period.label}  ·  환율·유가 평균 {fx_r:,.0f}원 / {fu_r:,.1f} USC', '이 시트의 월별 행이 RAW DATA 입니다. 파란 칸(CFG·운항횟수·L/F·A/R)을 고치면 이 시트 합계와 요약 시트가 함께 바뀝니다', "달마다 쓴 값의 출처는 아래 '근거' 를 보세요 (확정실적 = 경영기획 마감치 / 추정실적 = 주차별 보고 / CASK = 비용파일 추정)"] + [f"{d['y'] % 100:02d}.{d['m']:02d} : {d['src']}" for d in got])
    scenarios = [Scenario('실적 기준', results=period_results)]
    monthly = [Scenario('실적 기준', results=monthly_results)]
COMBI_COLS = [('년월', 10), ('기종', 16), ('CFG', 7), ('운항횟수\n(왕복)', 9), ('편수\n(OW)', 8), ('공급석', 10), ('수송석', 10), ('1왕복수입\nL/F(%)', 9), ('1왕복수입\nA/R(원)', 10), ('1왕복수입\n여객', 12), ('1왕복수입\n부대', 11), ('1왕복수입\n화물', 11), ('1왕복수입\n총수입', 12), ('1왕복수지\n변동비', 12), ('1왕복수지\n총비용', 12), ('1왕복수지\n한계이익', 12), ('1왕복수지\n영업이익', 12), ('1왕복수지\n영업이익률', 10), ('총수지\n총수입', 13), ('총수지\n총비용', 13), ('총수지\n영업이익', 13), ('환율(원)', 9), ('유가(USC)', 9), ('근거', 46)]
RT_MONEY_I = (9, 10, 11, 12, 13, 14, 15)
TOT_MONEY_I = (18, 19)
PROFIT_I = (16, 20)
PCT_I = (17,)

def _combined_tables(path, jobs, ds):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    used = set(wb.sheetnames)
    for route in dict.fromkeys((r for _, blocks in jobs for r, _ in blocks)):
        name = next((n for n in wb.sheetnames if n == _sheet_name(route, set())), None)
        if name is None:
            continue
        ws = wb[name]
        row = ws.max_row + 2
        for period, blocks in jobs:
            rows = next((rs for r, rs in blocks if r == route), [])
            got = [d for d in rows if not d['empty']]
            if not got:
                continue
            _text(ws, row, 2, f'■ 월별 수지 (기종 합계) - {period.label}   (금액 : 천원)', F_SUM)
            row += 1
            ws.row_dimensions[row].height = 32
            for i, (head, w) in enumerate(COMBI_COLS):
                c = ws.cell(row, 2 + i, head)
                c.font, c.fill, c.alignment, c.border = (F_HEAD_W, HEAD_FILL, CENTER, BOX)
                cur_w = ws.column_dimensions[L(2 + i)].width or 0
                ws.column_dimensions[L(2 + i)].width = max(cur_w, w)
            row += 1
            first = row
            for d in got:
                vals = _combi_vals(f"{d['y'] % 100}.{d['m']:02d}월", d['ac'], d) + [round(d.get('fx') or 0), round(d.get('fuel') or 0, 1), d['src']]
                _combi_row(ws, row, vals)
                row += 1
            tot_rt = max(sum((d['rt'] for d in got)), 1e-09)
            agg_d = {k: sum((d[k] for d in got)) for k in ('rt', 'ow', 'seats', 'pax', 'rev', 'anc', 'cargo', 'total_rev', 'cost', 'var')}
            agg_d['lf'] = agg_d['pax'] / agg_d['seats'] if agg_d['seats'] else 0
            agg_d['ar'] = agg_d['rev'] / agg_d['pax'] if agg_d['pax'] else 0
            wfx = sum(((d.get('fx') or 0) * d['rt'] for d in got)) / tot_rt
            wfu = sum(((d.get('fuel') or 0) * d['rt'] for d in got)) / tot_rt
            vals = _combi_vals('합계', '', agg_d) + [round(wfx), round(wfu, 1), '환율·유가는 운항횟수 가중평균']
            _combi_row(ws, row, vals, total=True)
            row += 3
    wb.save(path)
    return path

def _combi_vals(label, ac, d):
    rt = max(d['rt'], 1e-09)
    ow = d['ow']
    per = lambda k: d[k] / rt
    margin = d['total_rev'] - d['var']
    profit = d['total_rev'] - d['cost']
    return [label, ac, round(d['seats'] / ow) if ow else 0, round(d['rt'], 1), round(ow), round(d['seats']), round(d['pax']), d['lf'], round(d['ar']), round(per('rev')), round(per('anc')), round(per('cargo')), round(per('total_rev')), round(per('var')), round(per('cost')), round(margin / rt), round(profit / rt), profit / d['total_rev'] if d['total_rev'] else 0, round(d['total_rev']), round(d['cost']), round(profit)]

def _combi_row(ws, row, vals, total: bool=False):
    ws.row_dimensions[row].height = ROW_H
    for i, v in enumerate(vals):
        c = ws.cell(row, 2 + i, v)
        c.font = F_SUM if total else F_BODY
        c.border, c.alignment = (BOX, LEFT if i == len(vals) - 1 else CENTER)
        if total:
            c.fill = HEAD_FILL
            c.font = F_HEAD_W
        if i == 7:
            c.number_format = '#,##0.0%'
        elif i in PCT_I:
            c.number_format = '#,##0.0%' if total else PCT
        elif i in PROFIT_I:
            c.number_format = MONEY_K if total else PL_K
        elif i in RT_MONEY_I or i in TOT_MONEY_I:
            c.number_format = MONEY_K
        elif i in (3, 22):
            c.number_format = '#,##0.0'
        elif i >= 2 and isinstance(v, (int, float)):
            c.number_format = '#,##0'

def build_report(path, jobs, alloc, ds, sched, mg=None, wk=None):
    live = [d for _, blocks in jobs for _, rows in blocks for d in rows if not d['empty']]
    season_fallback = any((d.get('fallback') for d in live))
    pool_fix = any((d.get('pool_fix') and (not d.get('weekly')) for d in live))
    scenarios, monthly, route_assumptions = ([], [], {})
    for period, blocks in jobs:
        period_results, monthly_results = ([], [])
        _fill(period, blocks, period_results, monthly_results, route_assumptions, alloc, ds)
        scenarios.append(Scenario(period.label, results=period_results))
        monthly.append(Scenario(period.label, results=monthly_results))
    period = jobs[0][0]
    labels = ' · '.join((pp.label for pp, _ in jobs))
    ext = _month_ranges(live, '확정') or _month_ranges(live, '추정')
    assumptions = [f'기간 : {labels}'] + _cost_lines(live, alloc, mg, wk) + ['수입 : 확정·추정 달은 그 파일의 운송·부대·화물수입, 나머지 달은 과거실적 DATA 의 실제 운송·화물수입 (부대수입은 대노선별 RASK 기준 추정)', '운항횟수 : 확정·추정 달은 그 파일 편수, 나머지 달은 과거실적 DATA 의 실제 편수' + ('(실적이 없으면 확정 스케줄' + (f' · {sched}' if sched else '') + ')'), "조건 : 비교 없이 노선 자체의 월별 수지 · 달마다 쓴 값의 출처는 노선 시트 '근거' 에 적혀 있습니다"]
    rt_all = sum((d['rt'] for d in live))
    if rt_all:
        fx = sum(((d.get('fx') or 0) * d['rt'] for d in live)) / rt_all
        fu = sum(((d.get('fuel') or 0) * d['rt'] for d in live)) / rt_all
        assumptions.insert(1, f"비용 : 환율·유가 기간 평균 {fx:,.0f}원 / {fu:,.1f} USC (운항횟수 가중평균 · 달별 값은 노선 시트 '기종 합계' 표)")
    if season_fallback:
        assumptions.append('배부단가 : 대상 기간이 비용파일 시즌(W26) 밖이라 시즌 평균 배부단가를 대신 썼습니다 (환율·유가는 해당 월 INDEX 그대로)')
    if pool_fix:
        assumptions.append('고정비 POOL 이 일부 기간치뿐인 달(시즌 시작 월)은 배부단가만 가장 가까운 정상 달 것을 썼습니다 - 그대로 두면 고정비가 1/4 수준으로 빠집니다')
    narrative = build_narrative(scenarios, assumptions)
    out = write_excel(path, f'노선 수지 ({labels})', scenarios, assumptions, narrative, monthly=monthly, route_assumptions=route_assumptions, period_label=labels)
    return _combined_tables(out, jobs, ds)

def main():
    ap = argparse.ArgumentParser(description='노선 수지 (비교 없이 한 노선의 월별 손익)')
    ap.add_argument('routes', nargs='+', help='노선 (PUSKIX, PUS-KIX ...)')
    ap.add_argument('--period', default=['S26'], nargs='+', help='대상 기간. 2개까지 (예: --period S26 W26)')
    ap.add_argument('--sched', default=None, help='실적 없는 달에 쓸 운항스케줄 (DAILY, 2 DAILY, 4/W D1346 ...)')
    ap.add_argument('--ac', default=None, help='실적 없는 달에 쓸 기종 (기본: 그 노선 주력 기종)')
    ap.add_argument('--fixed-alloc', choices=('revenue', 'volume'), default='volume')
    ap.add_argument('--no-weekly', action='store_true', help='주차별 추정실적을 쓰지 않고 전 기간 CASK 로만 계산')
    ap.add_argument('--out', default=None)
    args = ap.parse_args()
    ds, act = (Dataset(), Actuals())
    wk = None if args.no_weekly else Weekly()
    mg = None if args.no_weekly else Mgmt()
    cb = None if args.no_weekly else CostBook()
    pl = Plan()
    labels = [x for raw in args.period for x in str(raw).split(',') if x.strip()]
    periods = []
    for lab in labels[:2]:
        try:
            periods.append(Period.parse(lab.strip()))
        except Exception:
            sys.exit(f'기간 인식 실패 : {lab}')
    if len(labels) > 2:
        eprint(f"   ! 기간은 2개까지만 씁니다 (뒤의 {', '.join(labels[2:])} 은 무시)")
    ac_hint = ds.resolve_aircraft(args.ac) if args.ac else None
    if args.ac and (not ac_hint):
        sys.exit(f'기종 인식 실패 : {args.ac}')
    jobs = []
    for period in periods:
        blocks = []
        for raw in args.routes:
            route = ds.resolve_route(raw)
            if not route:
                eprint(f'   ! 노선 인식 실패 : {raw}')
                continue
            rows = list(month_rows(ds, act, route, period, args.sched, ac_hint, args.fixed_alloc, item=len(blocks), wk=wk, mg=mg, cb=cb, pl=pl))
            blocks.append((route, rows))
            eprint(f'\n■ {route} ({period.label})')
            eprint(f"   {'년월':<8}{'편수':>6}{'L/F':>8}{'A/R':>9}{'총수입':>12}{'총비용':>12}{'영업이익':>12}{'이익률':>8}  근거")
            tot = dict(ow=0, rev=0.0, cost=0.0, seats=0.0, pax=0.0, anc=0.0, cargo=0.0)
            for d in rows:
                if d['empty']:
                    eprint(f"   {d['y'] % 100:02d}.{d['m']:02d}   {'-':>6}{'':>8}{'':>9}{'':>12}{'':>12}{'':>12}{'':>8}  {d['src']}")
                    continue
                p = d['total_rev'] - d['cost']
                eprint(f"   {d['y'] % 100:02d}.{d['m']:02d}{d['ow']:>7,}{d['lf']:>8.1%}{d['ar']:>9,.0f}{d['total_rev'] / EOK:>11,.1f}억{d['cost'] / EOK:>11,.1f}억{p / EOK:>11,.1f}억{p / d['total_rev']:>8.1%}  {d['src']}")
                for k, v in (('ow', d['ow']), ('rev', d['rev']), ('cost', d['cost']), ('seats', d['seats']), ('pax', d['pax']), ('anc', d['anc']), ('cargo', d['cargo'])):
                    tot[k] += v
            trev = tot['rev'] + tot['anc'] + tot['cargo']
            if tot['seats']:
                eprint(f"   {'합계':<6}{tot['ow']:>7,}{tot['pax'] / tot['seats']:>8.1%}{tot['rev'] / tot['pax']:>9,.0f}{trev / EOK:>11,.1f}억{tot['cost'] / EOK:>11,.1f}억{(trev - tot['cost']) / EOK:>11,.1f}억{(trev - tot['cost']) / trev:>8.1%}")
        if blocks:
            jobs.append((period, blocks))
    if not jobs:
        sys.exit('계산할 노선이 없습니다.')
    tag = '+'.join((r.replace(' V.V', '').replace('-', '') for r, _ in jobs[0][1]))[:50]
    span = '+'.join((pp.label for pp, _ in jobs))
    out = args.out or BASE / 'output' / f'노선수지_{datetime.now():%y%m%d} ({tag} {span}).xlsx'
    __import__('pathlib').Path(out).parent.mkdir(parents=True, exist_ok=True)
    build_report(out, jobs, args.fixed_alloc, ds, args.sched, mg, wk)
    eprint(f'\n[저장] {out}')
if __name__ == '__main__':
    main()
